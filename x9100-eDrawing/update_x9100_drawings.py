#!/usr/bin/env python3
"""
========================================================================================
Project: KALASIN STARCH PLANT - JET COOKER PROJECT
Folder:  E:\\xApp-01\\x260911-002-KalasinEngineering\\x9100-eDrawing
Script:  update_x9100_drawings.py
Purpose: Copy and update all DI (1756-IB32), DO (1756-OB32), and AI (1756-IF16) CAD DXF
         drawings using data from 'IO_List-By_SlotConfig_rev02.xlsx'.
         Updates PLC Wire Tags, Terminal Block Tags, Block Attributes, Header Text,
         Location Headers, and Relay/Analog Channel Configurations.
========================================================================================
"""

import os
import sys
import openpyxl
import ezdxf

sys.stdout.reconfigure(encoding='utf-8')

def safe_saveas(doc, filepath):
    try:
        doc.saveas(filepath)
        return True
    except PermissionError:
        print(f"  [Warning] Permission denied writing to {filepath} (file may be open in AutoCAD). Skipping file.")
        return False
    except Exception as e:
        print(f"  [Error] Failed saving {filepath}: {e}")
        return False

def main():
    base_dir = os.path.abspath(os.path.dirname(__file__))
    excel_path = os.path.join(base_dir, "IO_List-By_SlotConfig_rev02.xlsx")
    template_di = os.path.join(base_dir, "template_DI-r02.dxf")
    template_do = os.path.join(base_dir, "template_DO-r02.dxf")
    template_ai = os.path.join(base_dir, "template_AI-r02.dxf")

    if not os.path.exists(excel_path):
        print(f"Error: Excel file not found at {excel_path}")
        return

    print(f"Loading Excel workbook: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    target_di = [
        "C1S4", "C1S5", "C1S6", "C1S7",
        "C2S1", "C2S2", "C2S3", "C2S4", "C2S5",
        "C3S1", "C3S2", "C3S3",
        "C4S1", "C4S2", "C4S3",
        "C5S1", "C5S2"
    ]

    target_do = [
        "C1S8", "C1S9",
        "C2S6", "C2S7", "C2S8",
        "C3S4", "C3S5",
        "C4S4",
        "C5S3"
    ]

    target_ai = [
        "C1S10", "C1S11",
        "C2S9", "C2S10", "C2S11", "C2S12",
        "C3S6", "C3S7", "C3S8", "C3S9",
        "C4S5", "C4S6",
        "C5S4"
    ]

    even_pins = [2, 4, 6, 8, 10, 12, 14, 16, 18, 20, 22, 24, 26, 28, 30, 32, 34, 36]
    odd_pins = [1, 3, 5, 7, 9, 11, 13, 15, 17, 19, 21, 23, 25, 27, 29, 31, 33, 35]

    term_to_pin = {}
    for k in range(1, 17):
        term_to_pin[k] = k
    for k in range(17, 33):
        term_to_pin[k] = k + 2

    ai_term_to_pin = {
        1: 2,   2: 1,   # 1A (IN-0), 1B (i RTN-0)
        3: 4,   4: 3,   # 2A (IN-1), 2B (i RTN-1)
        5: 6,   6: 5,   # 3A (IN-2), 3B (i RTN-2)
        7: 8,   8: 7,   # 4A (IN-3), 4B (i RTN-3)
        9: 12,  10: 11, # 5A (IN-4), 5B (i RTN-4)
        11: 14, 12: 13, # 6A (IN-5), 6B (i RTN-5)
        13: 16, 14: 15, # 7A (IN-6), 7B (i RTN-6)
        15: 18, 16: 17, # 8A (IN-7), 8B (i RTN-7)
        17: 20, 18: 19, # 9A (IN-8), 9B (i RTN-8)
        19: 22, 20: 21, # 10A (IN-9), 10B (i RTN-9)
        21: 24, 22: 23, # 11A (IN-10), 11B (i RTN-10)
        23: 26, 24: 25, # 12A (IN-11), 12B (i RTN-11)
        25: 30, 26: 29, # 13A (IN-12), 13B (i RTN-12)
        27: 32, 28: 31, # 14A (IN-13), 14B (i RTN-13)
        29: 34, 30: 33, # 15A (IN-14), 15B (i RTN-14)
        31: 36, 32: 35  # 16A (IN-15), 16B (i RTN-15)
    }

    def parse_slot_data(sheet):
        pin_data = {}
        tb_name = ""
        active_count = 0
        spare_count = 0

        rack_meta = str(sheet.cell(row=3, column=1).value or "")
        dest_area = "CA1 / FIELD"
        if "DESTINATION AREA:" in rack_meta:
            dest_area = rack_meta.split("DESTINATION AREA:")[1].strip()

        for r in range(6, 42):
            pin_num = sheet.cell(r, 1).value
            if pin_num is None: continue
            pin_num = int(pin_num)
            term_desc = str(sheet.cell(r, 2).value or "")
            tb_col = str(sheet.cell(r, 3).value or "")
            plc_tag_side = str(sheet.cell(r, 4).value or "")
            term_tag_side = str(sheet.cell(r, 5).value or "")
            dest = str(sheet.cell(r, 6).value or "")
            plc_tag = str(sheet.cell(r, 7).value or "")
            inst_tag = str(sheet.cell(r, 8).value or "")
            inst_desc = str(sheet.cell(r, 9).value or "")
            status = str(sheet.cell(r, 10).value or "").upper()

            if "-" in tb_col and tb_col not in ("0VDC", "+24VDC") and not tb_name:
                tb_name = tb_col.rsplit("-", 1)[0]
                if "(" in tb_name:
                    tb_name = tb_name.rsplit("(", 1)[0]

            if status == "ACTIVE": active_count += 1
            elif status == "SPARE": spare_count += 1

            pin_data[pin_num] = {
                "desc": term_desc, "tb_col": tb_col, "plc_tag_side": plc_tag_side,
                "term_tag_side": term_tag_side, "dest": dest, "plc_tag": plc_tag,
                "inst_tag": inst_tag, "inst_desc": inst_desc, "status": status
            }

        return pin_data, tb_name, dest_area, active_count, spare_count

    def add_io_list_table(msp, slot_key, pin_data, module_type="DI"):
        # Bounding box matching the new template I/O List frame (X: 65.7..1411.5, Y: 333.8..965.9)
        X_START = 65.7
        X_END = 1411.5
        Y_TOP = 965.9
        Y_BOTTOM = 333.8
        
        HEADER_HEIGHT = 28.0
        ROW_HEIGHT = (Y_TOP - HEADER_HEIGHT - Y_BOTTOM) / 32.0  # ~18.88 mm per row
        
        col_x = [65.7, 125.0, 200.0, 400.0, 580.0, 1210.0, 1411.5]
        headers = ["PIN", "TERM", "PLC TAG", "INST TAG", "SIGNAL DESCRIPTION", "DEST / JB"]
        
        # Outer Frame Box
        msp.add_lwpolyline(
            [(X_START, Y_BOTTOM), (X_END, Y_BOTTOM), (X_END, Y_TOP), (X_START, Y_TOP), (X_START, Y_BOTTOM)],
            dxfattribs={'layer': '0', 'color': 7}
        )
        y_hdr = Y_TOP - HEADER_HEIGHT
        msp.add_line((X_START, y_hdr), (X_END, y_hdr), dxfattribs={'layer': '0', 'color': 7})
        
        # Column Dividers
        for x in col_x[1:-1]:
            msp.add_line((x, Y_BOTTOM), (x, Y_TOP), dxfattribs={'layer': '0', 'color': 7})
            
        # Table Title Header (Update existing 'I/O List' text or add new text)
        found_title = False
        for e in msp:
            if e.dxftype() in ('TEXT', 'MTEXT'):
                txt = e.dxf.text if hasattr(e.dxf, 'text') else e.text
                if txt == 'I/O List':
                    if hasattr(e.dxf, 'text'):
                        e.dxf.text = f"SLOT {slot_key} - I/O POINT SCHEDULE & CHANNEL LIST"
                    else:
                        e.text = f"SLOT {slot_key} - I/O POINT SCHEDULE & CHANNEL LIST"
                    found_title = True
                    
        if not found_title:
            msp.add_text(
                f"SLOT {slot_key} - I/O POINT SCHEDULE & CHANNEL LIST",
                dxfattribs={'height': 11.0, 'insert': (X_START + 8.0, Y_TOP + 6.0), 'layer': '0', 'color': 2}
            )
        
        # Column Header Texts
        for i, h_text in enumerate(headers):
            cy = Y_TOP - (HEADER_HEIGHT / 2.0) - 3.0
            msp.add_text(
                h_text,
                dxfattribs={'height': 7.5, 'insert': (col_x[i] + 3.0, cy), 'layer': '0', 'color': 3}
            )
            
        term_to_pin_di_do = {k: (k if k <= 16 else k + 2) for k in range(1, 33)}
        ai_term_to_pin_map = {
            1: 2, 2: 1, 3: 4, 4: 3, 5: 6, 6: 5, 7: 8, 8: 7,
            9: 12, 10: 11, 11: 14, 12: 13, 13: 16, 14: 15, 15: 18, 16: 17,
            17: 20, 18: 19, 19: 22, 20: 21, 21: 24, 22: 23, 23: 26, 24: 25,
            25: 30, 26: 29, 27: 32, 28: 31, 29: 34, 30: 33, 31: 36, 32: 35
        }
        
        pin_map = ai_term_to_pin_map if module_type == "AI" else term_to_pin_di_do

        for r_idx in range(1, 33):
            y_row_bot = y_hdr - r_idx * ROW_HEIGHT
            y_text = y_row_bot + 3.0
            
            if r_idx < 32:
                msp.add_line((X_START, y_row_bot), (X_END, y_row_bot), dxfattribs={'layer': '0', 'color': 8})
                
            pin = pin_map.get(r_idx, r_idx)
            p_info = pin_data.get(pin, {})
            
            pin_str = str(pin)
            term_str = p_info.get("desc", f"Pin {pin}")
            inst_tag = p_info.get("inst_tag", "")
            inst_desc = p_info.get("inst_desc", "")
            dest_jb = p_info.get("dest", "")
            
            clean_inst_tag = inst_tag.strip() if inst_tag and inst_tag.strip() not in ("-", "Spare", "None", "0", "0.0", "nan") else "Spare"

            if module_type == "AI":
                ch = (r_idx - 1) // 2
            else:
                ch = r_idx - 1
            ch_tag = f"{ch:02d}"

            plc_tag = f"{module_type}_{slot_key}_{ch_tag}_{clean_inst_tag}"

            if inst_tag in ("-", "Spare", "None", "0"): inst_tag = ""
            if inst_desc in ("-", "Spare", "None", "0"): inst_desc = ""
            if dest_jb in ("-", "Spare", "None", "0"): dest_jb = ""
            
            msp.add_text(pin_str, dxfattribs={'height': 5.5, 'insert': (col_x[0] + 3.0, y_text), 'layer': '0', 'color': 7})
            msp.add_text(term_str[:10], dxfattribs={'height': 5.5, 'insert': (col_x[1] + 3.0, y_text), 'layer': '0', 'color': 7})
            msp.add_text(plc_tag[:30], dxfattribs={'height': 5.5, 'insert': (col_x[2] + 3.0, y_text), 'layer': '0', 'color': 7})
            msp.add_text(inst_tag[:16], dxfattribs={'height': 5.5, 'insert': (col_x[3] + 3.0, y_text), 'layer': '0', 'color': 7})
            msp.add_text(inst_desc[:55], dxfattribs={'height': 5.5, 'insert': (col_x[4] + 3.0, y_text), 'layer': '0', 'color': 7})
            msp.add_text(dest_jb[:20], dxfattribs={'height': 5.5, 'insert': (col_x[5] + 3.0, y_text), 'layer': '0', 'color': 7})

    generated_di = 0
    generated_do = 0
    generated_ai = 0

    print("\n=======================================================")
    print("PROCESSING DIGITAL INPUT (DI) SLOTS (1756-IB32)")
    print("=======================================================")

    for slot_key in target_di:
        if slot_key not in wb.sheetnames:
            print(f"Warning: Sheet {slot_key} not in workbook, skipping.")
            continue

        sheet = wb[slot_key]
        chassis_str = slot_key[:2]
        slot_num_str = slot_key[3:]
        chassis_num = slot_key[1]

        pin_data, tb_name, dest_area, active_count, spare_count = parse_slot_data(sheet)
        if not tb_name:
            tb_name = f"P{chassis_num}-TBDI{slot_num_str}"

        print(f"[{slot_key}] Generating DI DXF (TB: {tb_name}, Active: {active_count}, Spare: {spare_count})...")

        doc = ezdxf.readfile(template_di)
        msp = doc.modelspace()

        # 1. Update Header Title & Block Attributes
        slot_title = f"CHASSIS {chassis_str} SLOT {slot_num_str}"
        for t in msp:
            if t.dxftype() in ('TEXT', 'ATTDEF') and (hasattr(t.dxf, 'tag') and t.dxf.tag == 'B' or 'CHASSIS' in getattr(t.dxf, 'text', '')):
                t.dxf.text = slot_title

        for ins in msp:
            if ins.dxftype() == 'INSERT' and ins.dxf.name == 'TPA':
                for a in ins.attribs:
                    if a.dxf.tag == 'B':
                        a.dxf.text = slot_title

        # 2. Location / Cabinet headers
        for t in msp:
            if t.dxftype() == 'TEXT':
                if t.dxf.text.startswith("=CA"):
                    t.dxf.text = f"=CA{chassis_num}+MCP-M{chassis_num}-P{slot_num_str}"
                elif "=JB" in t.dxf.text or "=Field" in t.dxf.text:
                    if dest_area and dest_area != "CA1":
                        first_jb = dest_area.split(",")[0].strip()
                        t.dxf.text = f"={first_jb}"

        # 3. Terminal Block Header Text
        for t in msp:
            if t.dxftype() == 'TEXT' and ('P1-TBDI1' in t.dxf.text or 'P1-TBRY1' in t.dxf.text or 'TBDI' in t.dxf.text):
                t.dxf.text = tb_name

        # 4. Left PLC Card Wire Tags (Even pins) at X ~ 365.8
        left_plc = [t for t in msp if t.dxftype() == 'TEXT' and 340.0 < t.dxf.insert[0] < 400.0 and t.dxf.insert[1] > 1000.0]
        left_plc.sort(key=lambda t: -t.dxf.insert[1])
        if len(left_plc) == 18:
            for pin, t in zip(even_pins, left_plc):
                if pin in pin_data:
                    t.dxf.text = pin_data[pin]["plc_tag_side"]

        # 5. Right PLC Card Wire Tags (Odd pins) at X ~ 1048.3
        right_plc = [t for t in msp if t.dxftype() == 'TEXT' and 1020.0 < t.dxf.insert[0] < 1080.0 and t.dxf.insert[1] > 1000.0]
        right_plc.sort(key=lambda t: -t.dxf.insert[1])
        if len(right_plc) == 18:
            for pin, t in zip(odd_pins, right_plc):
                if pin in pin_data:
                    t.dxf.text = pin_data[pin]["plc_tag_side"]

        # 6. Terminal Block Left Wire Tags (PLC Tag Side at TB) at X ~ 1447.9
        tb_left = [t for t in msp if t.dxftype() == 'TEXT' and 1400.0 < t.dxf.insert[0] < 1500.0 and t.dxf.insert[1] > 1000.0]
        tb_left.sort(key=lambda t: -t.dxf.insert[1])
        if len(tb_left) == 32:
            for k in range(1, 33):
                pin = term_to_pin[k]
                if pin in pin_data:
                    tb_left[k-1].dxf.text = pin_data[pin]["plc_tag_side"]

        # 7. Terminal Block Right Wire Tags (Terminal Tag Side at TB) at X ~ 1789.9
        tb_right = [t for t in msp if t.dxftype() == 'TEXT' and 1750.0 < t.dxf.insert[0] < 1830.0 and t.dxf.insert[1] > 1000.0]
        tb_right.sort(key=lambda t: -t.dxf.insert[1])
        if len(tb_right) == 32:
            for k in range(1, 33):
                pin = term_to_pin[k]
                if pin in pin_data:
                    tb_right[k-1].dxf.text = pin_data[pin]["term_tag_side"]

        # 8. Destination / Field Wire Tags at X ~ 2353.2
        dest_tags = [t for t in msp if t.dxftype() == 'TEXT' and 2300.0 < t.dxf.insert[0] < 2400.0 and t.dxf.insert[1] < 2450.0]
        dest_tags.sort(key=lambda t: -t.dxf.insert[1])
        if len(dest_tags) == 32:
            for k in range(1, 33):
                pin = term_to_pin[k]
                if pin in pin_data:
                    dest_tags[k-1].dxf.text = pin_data[pin]["term_tag_side"]

        # 9. Add I/O List Frame Table
        add_io_list_table(msp, slot_key, pin_data, "DI")

        ch_dir = os.path.join(base_dir, chassis_str)
        os.makedirs(ch_dir, exist_ok=True)
        out_file1 = os.path.join(ch_dir, f"{slot_key}-IB32_Wiring_Diagram.dxf")
        out_file2 = os.path.join(ch_dir, f"{slot_key}-IB32.dxf")
        s1 = safe_saveas(doc, out_file1)
        s2 = safe_saveas(doc, out_file2)
        if s1 or s2:
            generated_di += 1
            print(f"  -> Saved {slot_key} DI DXF")

    print("\n=======================================================")
    print("PROCESSING DIGITAL OUTPUT (DO) SLOTS (1756-OB32)")
    print("=======================================================")

    for slot_key in target_do:
        if slot_key not in wb.sheetnames:
            print(f"Warning: Sheet {slot_key} not in workbook, skipping.")
            continue

        sheet = wb[slot_key]
        chassis_str = slot_key[:2]
        slot_num_str = slot_key[3:]
        chassis_num = slot_key[1]

        pin_data, tb_name, dest_area, active_count, spare_count = parse_slot_data(sheet)
        if not tb_name:
            tb_name = f"P{chassis_num}-TBRL1"

        print(f"[{slot_key}] Generating DO DXF (TB: {tb_name}, Active: {active_count}, Spare: {spare_count})...")

        doc = ezdxf.readfile(template_do)
        msp = doc.modelspace()

        # 1. Update Header Title & Module Type & Attributes
        slot_title = f"CHASSIS {chassis_str} SLOT {slot_num_str}"
        for t in msp:
            if t.dxftype() in ('TEXT', 'ATTDEF'):
                if hasattr(t.dxf, 'tag') and t.dxf.tag == 'B' or 'CHASSIS' in getattr(t.dxf, 'text', ''):
                    t.dxf.text = slot_title
                elif getattr(t.dxf, 'text', '') == '1756-IB32':
                    t.dxf.text = '1756-OB32'
                elif getattr(t.dxf, 'text', '') == 'DIGITAL INPUT MODULE':
                    t.dxf.text = 'DIGITAL OUTPUT MODULE'

        for ins in msp:
            if ins.dxftype() == 'INSERT' and ins.dxf.name == 'TPA':
                for a in ins.attribs:
                    if a.dxf.tag == 'B':
                        a.dxf.text = slot_title

        # 2. Location / Cabinet headers
        for t in msp:
            if t.dxftype() == 'TEXT':
                if t.dxf.text.startswith("=CA"):
                    t.dxf.text = f"=CA{chassis_num}+MCP-M{chassis_num}-P{slot_num_str}"
                elif "=JB" in t.dxf.text or "=Field" in t.dxf.text:
                    if dest_area and dest_area != "CA1":
                        first_jb = dest_area.split(",")[0].strip()
                        t.dxf.text = f"={first_jb}"

        # 3. Terminal Block Header Text
        for t in msp:
            if t.dxftype() == 'TEXT' and ('P1-TBDI1' in t.dxf.text or 'P1-TBRY1' in t.dxf.text or 'P1-TBRL1' in t.dxf.text or 'TBRL' in t.dxf.text or 'TBRY' in t.dxf.text or 'TBDI' in t.dxf.text):
                t.dxf.text = tb_name

        # 4. Left PLC Card Wire Tags (Even pins) at X ~ 365.8
        left_plc = [t for t in msp if t.dxftype() == 'TEXT' and 340.0 < t.dxf.insert[0] < 400.0 and t.dxf.insert[1] > 1000.0]
        left_plc.sort(key=lambda t: -t.dxf.insert[1])
        if len(left_plc) == 18:
            for pin, t in zip(even_pins, left_plc):
                if pin in pin_data:
                    t.dxf.text = pin_data[pin]["term_tag_side"]

        # 5. Right PLC Card Wire Tags (Odd pins) at X ~ 1048.3
        right_plc = [t for t in msp if t.dxftype() == 'TEXT' and 1020.0 < t.dxf.insert[0] < 1080.0 and t.dxf.insert[1] > 1000.0]
        right_plc.sort(key=lambda t: -t.dxf.insert[1])
        if len(right_plc) == 18:
            for pin, t in zip(odd_pins, right_plc):
                if pin in pin_data:
                    t.dxf.text = pin_data[pin]["term_tag_side"]

        # 6. Terminal Block Left Wire Tags (Terminal Tag Side at Relay/TB) at X ~ 2651.6
        tb_left = [t for t in msp if t.dxftype() == 'TEXT' and 2600.0 < t.dxf.insert[0] < 2700.0 and t.dxf.insert[1] > 1000.0]
        tb_left.sort(key=lambda t: -t.dxf.insert[1])
        if len(tb_left) == 32:
            for k in range(1, 33):
                pin = term_to_pin[k]
                if pin in pin_data:
                    tb_left[k-1].dxf.text = pin_data[pin]["term_tag_side"]

        # 7. Terminal Block Right Wire Tags (Terminal Tag Side at Relay/TB) at X ~ 2993.7
        tb_right = [t for t in msp if t.dxftype() == 'TEXT' and 2950.0 < t.dxf.insert[0] < 3050.0 and t.dxf.insert[1] > 1000.0]
        tb_right.sort(key=lambda t: -t.dxf.insert[1])
        if len(tb_right) == 32:
            for k in range(1, 33):
                pin = term_to_pin[k]
                if pin in pin_data:
                    tb_right[k-1].dxf.text = pin_data[pin]["term_tag_side"]

        # 8. SPDT Relay Running Numbers (1..32) at X ~ 1742.7
        relay_num_texts = [e for e in msp if e.dxftype() == 'TEXT' and 1700.0 < e.dxf.insert[0] < 1760.0 and 1000.0 < e.dxf.insert[1] < 2430.0]
        relay_num_texts.sort(key=lambda e: -e.dxf.insert[1])
        if len(relay_num_texts) == 32:
            for idx, t in enumerate(relay_num_texts, start=1):
                t.dxf.text = str(idx)

        # 8. Relay Block Attributes (xDO_3_Term or xDI_3_Term inserts)
        do_inserts = [e for e in msp if e.dxftype() == "INSERT" and e.dxf.name in ("xDO_3_Term", "xDI_3_Term")]
        do_inserts.sort(key=lambda e: -e.dxf.insert[1])
        for i, ins in enumerate(do_inserts):
            k = i + 1
            pin = term_to_pin.get(k)
            p_info = pin_data.get(pin, {})
            tag_name = p_info.get("inst_tag", f"DO_{k-1}")
            if not tag_name or tag_name in ("-", "Spare"):
                tag_name = f"SPARE_{k-1}"
            desc_text = p_info.get("inst_desc", f"Digital Output Channel {k-1}")
            if desc_text in ("0", "-", "None"):
                desc_text = f"DO Channel {k-1} Relay Circuit"

            for attrib in ins.attribs:
                if attrib.dxf.tag == "PLC_POINT":
                    attrib.dxf.text = f"{slot_key}-{pin}"
                elif attrib.dxf.tag == "TAG":
                    attrib.dxf.text = str(tag_name)
                elif attrib.dxf.tag == "Description":
                    attrib.dxf.text = str(desc_text)[:28]

        # 9. Add I/O List Frame Table
        add_io_list_table(msp, slot_key, pin_data, "DO")

        ch_dir = os.path.join(base_dir, chassis_str)
        os.makedirs(ch_dir, exist_ok=True)
        out_file1 = os.path.join(ch_dir, f"{slot_key}-OB32_Wiring_Diagram.dxf")
        out_file2 = os.path.join(ch_dir, f"{slot_key}-OB32.dxf")
        s1 = safe_saveas(doc, out_file1)
        s2 = safe_saveas(doc, out_file2)
        if s1 or s2:
            generated_do += 1
            print(f"  -> Saved {slot_key} DO DXF")

    print("\n=======================================================")
    print("PROCESSING ANALOG INPUT (AI) SLOTS (1756-IF16)")
    print("=======================================================")

    for slot_key in target_ai:
        if slot_key not in wb.sheetnames:
            print(f"Warning: Sheet {slot_key} not in workbook, skipping.")
            continue

        sheet = wb[slot_key]
        chassis_str = slot_key[:2]
        slot_num_str = slot_key[3:]
        chassis_num = slot_key[1]

        pin_data, tb_name, dest_area, active_count, spare_count = parse_slot_data(sheet)
        if not tb_name:
            tb_name = f"P{chassis_num}-TBAI{slot_num_str}"

        print(f"[{slot_key}] Generating AI DXF (TB: {tb_name}, Active: {active_count}, Spare: {spare_count})...")

        doc = ezdxf.readfile(template_ai)
        msp = doc.modelspace()

        # 1. Update Header Title & Block Attributes
        slot_title = f"CHASSIS {chassis_str} SLOT {slot_num_str}"
        for t in msp:
            if t.dxftype() in ('TEXT', 'ATTDEF') and (hasattr(t.dxf, 'tag') and t.dxf.tag == 'B' or 'CHASSIS' in getattr(t.dxf, 'text', '')):
                t.dxf.text = slot_title

        for ins in msp:
            if ins.dxftype() == 'INSERT' and ins.dxf.name == 'TPA':
                for a in ins.attribs:
                    if a.dxf.tag == 'B':
                        a.dxf.text = slot_title

        # 2. Location / Cabinet headers
        for t in msp:
            if t.dxftype() == 'TEXT':
                if t.dxf.text.startswith("=CA"):
                    t.dxf.text = f"=CA{chassis_num}+MCP-M{chassis_num}-P{slot_num_str}"
                elif "=JB" in t.dxf.text or "=Field" in t.dxf.text:
                    if dest_area and dest_area != "CA1":
                        first_jb = dest_area.split(",")[0].strip()
                        t.dxf.text = f"={first_jb}"

        # 3. Terminal Block Header Text
        for t in msp:
            if t.dxftype() == 'TEXT' and ('P1-TBDI1' in t.dxf.text or 'P1-TBAI1' in t.dxf.text or 'TBAI' in t.dxf.text or 'TBDI' in t.dxf.text):
                t.dxf.text = tb_name

        # 4. Left PLC Card Wire Tags (Even pins) at X ~ 365.8
        left_plc = [t for t in msp if t.dxftype() == 'TEXT' and 340.0 < t.dxf.insert[0] < 400.0 and t.dxf.insert[1] > 1000.0]
        left_plc.sort(key=lambda t: -t.dxf.insert[1])
        if len(left_plc) == 18:
            for pin, t in zip(even_pins, left_plc):
                if pin in pin_data:
                    t.dxf.text = pin_data[pin]["plc_tag_side"]

        # 5. Right PLC Card Wire Tags (Odd pins) at X ~ 1048.3
        right_plc = [t for t in msp if t.dxftype() == 'TEXT' and 1020.0 < t.dxf.insert[0] < 1080.0 and t.dxf.insert[1] > 1000.0]
        right_plc.sort(key=lambda t: -t.dxf.insert[1])
        if len(right_plc) == 18:
            for pin, t in zip(odd_pins, right_plc):
                if pin in pin_data:
                    t.dxf.text = pin_data[pin]["plc_tag_side"]

        # 6. Terminal Block Left Wire Tags (PLC Tag Side at TB) at X ~ 1400..1500
        tb_left = [t for t in msp if t.dxftype() == 'TEXT' and 1400.0 < t.dxf.insert[0] < 1500.0 and t.dxf.insert[1] > 1000.0]
        tb_left.sort(key=lambda t: -t.dxf.insert[1])
        if len(tb_left) == 32:
            for k in range(1, 33):
                pin = ai_term_to_pin[k]
                if pin in pin_data:
                    tb_left[k-1].dxf.text = pin_data[pin]["plc_tag_side"]

        # 7. Terminal Block Right Wire Tags (Terminal Tag Side at TB) at X ~ 1789.9
        tb_right = [t for t in msp if t.dxftype() == 'TEXT' and 1750.0 < t.dxf.insert[0] < 1830.0 and t.dxf.insert[1] > 1000.0]
        tb_right.sort(key=lambda t: -t.dxf.insert[1])
        if len(tb_right) == 32:
            for k in range(1, 33):
                pin = ai_term_to_pin[k]
                if pin in pin_data:
                    tb_right[k-1].dxf.text = pin_data[pin]["term_tag_side"]

        # 8. Destination / Field Wire Tags at X ~ 2353.2
        dest_tags = [t for t in msp if t.dxftype() == 'TEXT' and 2300.0 < t.dxf.insert[0] < 2400.0 and t.dxf.insert[1] < 2450.0]
        dest_tags.sort(key=lambda t: -t.dxf.insert[1])
        if len(dest_tags) == 32:
            for k in range(1, 33):
                pin = ai_term_to_pin[k]
                if pin in pin_data:
                    dest_tags[k-1].dxf.text = pin_data[pin]["term_tag_side"]

        # 9. Add I/O List Frame Table
        add_io_list_table(msp, slot_key, pin_data, "AI")

        ch_dir = os.path.join(base_dir, chassis_str)
        os.makedirs(ch_dir, exist_ok=True)
        out_file1 = os.path.join(ch_dir, f"{slot_key}-IF16_Wiring_Diagram.dxf")
        out_file2 = os.path.join(ch_dir, f"{slot_key}-IF16.dxf")
        s1 = safe_saveas(doc, out_file1)
        s2 = safe_saveas(doc, out_file2)
        if s1 or s2:
            generated_ai += 1
            print(f"  -> Saved {slot_key} AI DXF")

    print("\n=======================================================")
    print(f"Successfully generated/updated:")
    print(f"  - {generated_di} Digital Input (DI) slot drawings")
    print(f"  - {generated_do} Digital Output (DO) slot drawings")
    print(f"  - {generated_ai} Analog Input (AI) slot drawings")
    print(f"  - Total: {generated_di + generated_do + generated_ai} DXF drawings updated in {base_dir}")
    print("=======================================================")

if __name__ == '__main__':
    main()
