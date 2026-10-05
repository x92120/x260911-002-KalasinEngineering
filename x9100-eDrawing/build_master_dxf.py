#!/usr/bin/env python3
r"""
========================================================================================
Project: KALASIN STARCH PLANT - JET COOKER PROJECT
Folder:  E:\xApp-01\x260911-002-KalasinEngineering\x9100-eDrawing
Script:  build_master_dxf.py
Purpose: Create a single MASTER combined CAD DXF file containing ALL slot drawings
         arranged in a 5 Row x 13 Column grid layout (Chassis C1 to C5, Slots S0 to S12).
========================================================================================
"""

import os
import sys
import openpyxl
import ezdxf
from ezdxf.addons import Importer

sys.stdout.reconfigure(encoding='utf-8')

def safe_saveas(doc, filepath):
    try:
        doc.saveas(filepath)
        return True
    except PermissionError:
        print(f"  [Warning] Permission denied writing to {filepath} (file may be open in AutoCAD). Skipping file.")
        return False
    except Exception as e:
        print(f"  [Error] Failed saving {filepath}: {e}")
        return False

def create_placeholder_doc(template_path, slot_tag, chassis_str, slot_num_str, hw_card, sig_fam, dest_area):
    doc = ezdxf.readfile(template_path)
    msp = doc.modelspace()

    # 1. Update Title Header & Block Attributes
    for t in msp:
        if t.dxftype() == 'TEXT':
            if 'CHASSIS' in t.dxf.text:
                t.dxf.text = f"CHASSIS {chassis_str} SLOT {slot_num_str}"
            elif t.dxf.text in ('1756-IB32', '1756-OB32', '1756-IF16'):
                t.dxf.text = str(hw_card)
            elif 'MODULE' in t.dxf.text:
                t.dxf.text = str(sig_fam).upper()

    for ins in msp:
        if ins.dxftype() == 'INSERT' and ins.dxf.name == 'TPA':
            for a in ins.attribs:
                if a.dxf.tag == 'B':
                    a.dxf.text = f"CHASSIS {chassis_str} SLOT {slot_num_str}"

    # 2. Location / Cabinet headers
    for t in msp:
        if t.dxftype() == 'TEXT':
            if t.dxf.text.startswith("=CA"):
                t.dxf.text = f"=CA{chassis_str[1]}+MCP-M{chassis_str[1]}-P{slot_num_str}"
            elif "=JB" in t.dxf.text or "=Field" in t.dxf.text:
                if dest_area and dest_area != "CA1":
                    first_jb = dest_area.split(",")[0].strip()
                    t.dxf.text = f"={first_jb}"

    # 3. Clear wire tags in diagram area for non-wired slots
    for t in msp:
        if t.dxftype() == 'TEXT' and 340.0 < t.dxf.insert[0] < 2500.0 and t.dxf.insert[1] < 2450.0:
            if t.dxf.text.startswith("I-") or t.dxf.text.startswith("O-") or t.dxf.text.startswith("A-") or t.dxf.text.startswith("P1-") or t.dxf.text.startswith("DI_") or t.dxf.text.startswith("DO_") or t.dxf.text.startswith("AI_"):
                t.dxf.text = ""

    # 4. Add center description label inside the drawing box
    center_text = f"SLOT {slot_tag}\n{hw_card}\n{sig_fam}"
    msp.add_text(
        center_text,
        dxfattribs={
            'height': 32.0,
            'insert': (1400.0, 1500.0),
            'layer': '0',
            'color': 7
        }
    )
    return doc

def main():
    base_dir = os.path.abspath(os.path.dirname(__file__))
    excel_path = os.path.join(base_dir, "IO_List-By_SlotConfig_rev02.xlsx")
    template_di = os.path.join(base_dir, "template_DI-r02.dxf")

    if not os.path.exists(excel_path):
        print(f"Error: Excel file not found at {excel_path}")
        return

    print(f"Loading Excel slot directory from: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    index_sheet = wb['00_Slot_Index']

    slots_info = {}
    for r in range(14, index_sheet.max_row + 1):
        tag = index_sheet.cell(r, 1).value
        chassis = index_sheet.cell(r, 2).value
        slot_no = index_sheet.cell(r, 3).value
        hw_card = index_sheet.cell(r, 4).value
        sig_fam = index_sheet.cell(r, 5).value
        dest = index_sheet.cell(r, 7).value
        if tag:
            slots_info[str(tag).strip()] = {
                "chassis": str(chassis).strip(),
                "slot_no": slot_no,
                "hw_card": str(hw_card).strip(),
                "sig_fam": str(sig_fam).strip(),
                "dest": str(dest).strip()
            }

    master_doc = ezdxf.new("R2010")
    master_msp = master_doc.modelspace()

    # Define Grid Parameters with generous spacing between drawing frames
    # Frame size is ~3893mm Width x ~2704mm Height
    COL_PITCH = 5000.0  # Horizontal pitch (gives 1106mm gap between columns)
    ROW_PITCH = 3600.0  # Vertical pitch (gives 896mm gap between rows)
    
    chassis_list = ["C1", "C2", "C3", "C4", "C5"]
    chassis_names = {
        "C1": "CHASSIS C1 - MAIN CONTROLLER RACK (1756-A13 / CA1 MCP)",
        "C2": "CHASSIS C2 - MAIN EXPANSION I/O RACK (1756-A13 / CA1 MCP)",
        "C3": "CHASSIS C3 - REMOTE I/O RACK 1 (1756-A13 / SPRAY DRYER & PACKING TOWER)",
        "C4": "CHASSIS C4 - REMOTE I/O RACK 2 (1756-A13 / UPPER SPRAY DRYER 6TH/8TH FL)",
        "C5": "CHASSIS C5 - REMOTE I/O RACK 5 (1756-A7 / SLURRY BUILDING RIO-200)"
    }

    print("\n=======================================================")
    print("BUILDING MASTER COMBINED DXF (5 ROWS x 13 COLUMNS GRID)")
    print("=======================================================")

    # Add Overall Master Title Header at Top
    top_title_y = 4 * ROW_PITCH + 3200.0
    master_msp.add_text(
        "KALASIN STARCH PLANT - JET COOKER AUTOMATION SYSTEM (Ref: x2608003)",
        dxfattribs={'height': 100.0, 'insert': (0.0, top_title_y + 500.0), 'layer': '0', 'color': 2}
    )
    master_msp.add_text(
        "MASTER CONTROL SYSTEM & FIELD I/O ALL SLOTS WIRING DIAGRAM (CHASSIS C1 - C5)",
        dxfattribs={'height': 75.0, 'insert': (0.0, top_title_y + 300.0), 'layer': '0', 'color': 3}
    )
    master_msp.add_text(
        "LAYOUT: 5 CHASSIS RACKS x 13 SLOTS GRID (65 SLOT TILES)",
        dxfattribs={'height': 60.0, 'insert': (0.0, top_title_y + 140.0), 'layer': '0', 'color': 7}
    )

    total_inserted = 0

    for row_idx, ch_code in enumerate(chassis_list):
        row_y = (4 - row_idx) * ROW_PITCH
        ch_title = chassis_names.get(ch_code, f"CHASSIS {ch_code}")

        # Add Row Section Header
        master_msp.add_text(
            f"=== {ch_title} ===",
            dxfattribs={'height': 60.0, 'insert': (0.0, row_y + 2900.0), 'layer': '0', 'color': 1}
        )

        max_col_slots = 13 if ch_code != "C5" else 5

        for col_idx in range(13):
            slot_num = col_idx
            slot_tag = f"{ch_code}S{slot_num}"
            pos_x = col_idx * COL_PITCH
            pos_y = row_y

            ch_dir = os.path.join(base_dir, ch_code)
            dxf_filename = f"{slot_tag}.dxf"
            dxf_path = os.path.join(ch_dir, dxf_filename)
            if not os.path.exists(dxf_path) and os.path.exists(ch_dir):
                matching = [f for f in os.listdir(ch_dir) if f.startswith(f"{slot_tag}-") and f.endswith(".dxf") and "Wiring_Diagram" not in f]
                if matching:
                    dxf_path = os.path.join(ch_dir, matching[0])
                else:
                    dxf_path = os.path.join(base_dir, dxf_filename)

            s_info = slots_info.get(slot_tag, {})
            hw_card = s_info.get("hw_card", "1756-N2" if col_idx >= max_col_slots else "1756-CARD")
            sig_fam = s_info.get("sig_fam", "UNPOPULATED / RESERVE" if col_idx >= max_col_slots else "SPECIAL MODULE")
            dest_area = s_info.get("dest", "CA1")

            block_name = f"BLK_{slot_tag}"

            if os.path.exists(dxf_path):
                print(f"[{slot_tag}] Importing generated DXF drawing ({dxf_filename})...")
                src_doc = ezdxf.readfile(dxf_path)
            else:
                print(f"[{slot_tag}] Creating placeholder drawing for {hw_card} ({sig_fam})...")
                src_doc = create_placeholder_doc(template_di, slot_tag, ch_code, str(slot_num), hw_card, sig_fam, dest_area)

            importer = Importer(src_doc, master_doc)

            # Import required block definitions
            user_blocks = [b.name for b in src_doc.blocks if not b.name.startswith("*")]
            if user_blocks:
                importer.import_blocks(user_blocks)

            # Create target block and import modelspace
            target_blk = master_doc.blocks.new(name=block_name)
            importer.import_modelspace(target_blk)
            importer.finalize()

            # Insert block reference into master modelspace
            master_msp.add_blockref(block_name, (pos_x, pos_y))
            total_inserted += 1

    out_file1 = os.path.join(base_dir, "Master_All_Slots_Wiring_Diagram.dxf")
    out_file2 = os.path.join(base_dir, "Master_Slot_Config_All.dxf")

    s1 = safe_saveas(master_doc, out_file1)
    s2 = safe_saveas(master_doc, out_file2)

    print("\n=======================================================")
    print(f"Master DXF generation complete!")
    print(f"  - Inserted {total_inserted} slot blocks into grid (5 Rows x 13 Columns)")
    if s1: print(f"  - Output 1: {out_file1}")
    if s2: print(f"  - Output 2: {out_file2}")
    print("=======================================================")

if __name__ == '__main__':
    main()
