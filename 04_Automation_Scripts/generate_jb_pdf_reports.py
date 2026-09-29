#!/usr/bin/env python3
"""
Generate publication-quality engineering PDF reports for each Junction Box (and control enclosures)
from the master Excel file IO_List_xDev-R01-Tag35-5.xlsx.
"""

import os
import sys
import pandas as pd
import openpyxl
from fpdf import FPDF
from fpdf.enums import XPos, YPos

# Set output directories
WORKSPACE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = os.path.join(WORKSPACE_DIR, "reports_pdf")
os.makedirs(OUTPUT_DIR, exist_ok=True)

EXCEL_FILE = os.path.join(WORKSPACE_DIR, "IO_List_xDev-R01-Tag35-5.xlsx")

# --- Load Master Data ---
print(f"Loading data from {EXCEL_FILE}...")
xl = pd.ExcelFile(EXCEL_FILE)

# Load JB Config
df_jb_cfg_raw = xl.parse("Junction Box Config")
header_row = 0
for r_idx, row in df_jb_cfg_raw.iterrows():
    vals = [str(v).strip() for v in row if pd.notna(v)]
    if "Junction Name" in vals:
        header_row = r_idx
        break
df_jb_cfg = xl.parse("Junction Box Config", skiprows=header_row + 1)
df_jb_cfg.columns = [str(c).strip() for c in df_jb_cfg.columns]

jb_metadata = {}
for _, row in df_jb_cfg.iterrows():
    jb_name = str(row.get("Junction Name", "")).strip()
    if not jb_name or jb_name.lower() in ["nan", "total", "none"]:
        continue
    loc = str(row.get("Junction Loc", "")).strip() if pd.notna(row.get("Junction Loc")) else ""
    size = str(row.get("Size", "")).strip() if pd.notna(row.get("Size")) else ""
    is_ls = str(row.get("LS", "")).strip().upper() == "Y"
    is_jb = str(row.get("JB", "")).strip().upper() == "Y"
    
    if is_ls:
        jb_type = "Intrinsically Safe Junction Box (Ex i / LS)"
    elif jb_name.startswith("IS-"):
        jb_type = "Intrinsically Safe Junction Box (Ex i)"
    elif jb_name.startswith("RIO"):
        jb_type = "Remote I/O Enclosure (RIO Panel)"
    elif jb_name == "CA1":
        jb_type = "Main Control Cabinet (PLC CA1)"
    elif jb_name == "MCC":
        jb_type = "Motor Control Center Enclosure (MCC)"
    else:
        jb_type = "Field Junction Box (Standard)"
        
    jb_metadata[jb_name] = {
        "location": loc,
        "size": size,
        "type": jb_type
    }

if "CA1" in jb_metadata and not jb_metadata["CA1"]["size"]:
    jb_metadata["CA1"]["size"] = "1600-2000-800"
if "MCC" in jb_metadata:
    if not jb_metadata["MCC"]["location"]:
        jb_metadata["MCC"]["location"] = "MCC Room"
    if not jb_metadata["MCC"]["size"]:
        jb_metadata["MCC"]["size"] = "Draw-out Switchgear Panel"

# Load IO List
df_io = xl.parse("IO List")
df_io.columns = [str(c).strip() for c in df_io.columns]
# Exclude summary row (112) and empty destinations
df_io = df_io[df_io["Destination"].astype(str).str.strip() != "112"]
df_io = df_io[df_io["Destination"].notna()]

print(f"Total valid I/O channel records: {len(df_io)}")

def clean_val(val):
    if pd.isna(val):
        return "-"
    s = str(val).strip()
    if s.lower() in ["nan", "none", ""]:
        return "-"
    if s.endswith(".0") and s[:-2].isdigit():
        return s[:-2]
    return s

class IndustrialIOPDF(FPDF):
    def __init__(self, jb_name="All Junction Boxes", jb_info=None, is_master=False, stats=None):
        super().__init__(orientation="landscape", unit="mm", format="A4")
        self.jb_name = jb_name
        self.jb_info = jb_info or {}
        self.is_master = is_master
        self.stats = stats or {}
        self.is_summary_page = False
        self.set_auto_page_break(auto=True, margin=14)
        self.set_margins(10, 10, 10)
        self.alias_nb_pages()

    def header(self):
        # Cover page on master report has custom graphic layout
        if self.is_master and self.page_no() == 1:
            return

        # Top corporate banner
        self.set_font("Helvetica", "B", 10)
        self.set_fill_color(24, 43, 73) # Deep Navy
        self.set_text_color(255, 255, 255)
        self.cell(140, 6, "  INGREDION (THAILAND) CO., LTD.  |  PROJECT JET COOKER", 0, new_x=XPos.RIGHT, new_y=YPos.TOP, fill=True)
        self.cell(137, 6, "AEC INDUSTRIAL ENGINEERING  |  REV. 3.6  ", 0, new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="R", fill=True)
        
        # Subheader bar
        self.set_fill_color(240, 244, 248) # Light ice blue
        self.set_text_color(20, 30, 45)
        
        if self.is_summary_page:
            self.set_font("Helvetica", "B", 8.5)
            self.cell(277, 5.5, "  EXECUTIVE OVERVIEW  |  ALL ENCLOSURES SCHEDULE & CAPACITY MATRIX", 0, new_x=XPos.LMARGIN, new_y=YPos.NEXT, fill=True)
            self.ln(2)
            return

        self.set_font("Helvetica", "B", 8)
        loc_str = self.jb_info.get("location", "-")
        size_str = self.jb_info.get("size", "-")
        type_str = self.jb_info.get("type", "Junction Box")
        
        box_text = f"  ENCLOSURE: {self.jb_name}   |   TYPE: {type_str}   |   LOCATION: {loc_str}   |   SIZE: {size_str} mm"
        self.cell(277, 5, box_text, 0, new_x=XPos.LMARGIN, new_y=YPos.NEXT, fill=True)

        if self.stats:
            tot = self.stats.get("Total", 0)
            act = self.stats.get("Active", 0)
            sp = self.stats.get("Spare", 0)
            pct_sp = (sp / tot * 100) if tot > 0 else 0
            di = self.stats.get("DI", 0)
            do = self.stats.get("DO", 0)
            ai = self.stats.get("AI", 0)
            ao = self.stats.get("AO", 0)
            bus = self.stats.get("BUS", 0)
            
            self.set_font("Helvetica", "", 7.5)
            self.set_fill_color(225, 235, 245)
            stats_text = (f"  SUMMARY: Total Channels: {tot}   |   Active: {act}   |   "
                          f"Spare: {sp} ({pct_sp:.1f}%)   |   "
                          f"DI: {di}  |  DO: {do}  |  AI: {ai}  |  AO: {ao}  |  BUS: {bus}")
            self.cell(277, 4.5, stats_text, 0, new_x=XPos.LMARGIN, new_y=YPos.NEXT, fill=True)
            self.ln(2)
        else:
            self.ln(2)

    def footer(self):
        if self.is_master and self.page_no() == 1:
            return
        self.set_y(-10)
        self.set_font("Helvetica", "I", 7)
        self.set_text_color(100, 100, 100)
        
        doc_label = "Master Report (All Enclosures)" if self.is_master else self.jb_name
        self.cell(100, 5, f"Doc No.: Instrument I/O List (Kalasin) - {doc_label}", 0, new_x=XPos.RIGHT, new_y=YPos.TOP)
        self.cell(77, 5, f"Page {self.page_no()} of {{nb}}", 0, new_x=XPos.RIGHT, new_y=YPos.TOP, align="C")
        self.cell(100, 5, "CONFIDENTIAL  |  Generated: 2026-09-07", 0, new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="R")

COLUMNS = [
    {"id": "no", "name": "No.", "w": 8, "align": "C"},
    {"id": "terminal", "name": "Terminal", "w": 20, "align": "C"},
    {"id": "slot", "name": "Rack/Slot", "w": 16, "align": "C"},
    {"id": "channel", "name": "Channel", "w": 17, "align": "C"},
    {"id": "card", "name": "I/O Module", "w": 22, "align": "C"},
    {"id": "type", "name": "Type", "w": 11, "align": "C"},
    {"id": "signal", "name": "Signal Characteristic", "w": 25, "align": "L"},
    {"id": "tag", "name": "Instrument Tag", "w": 25, "align": "L"},
    {"id": "plc_tag", "name": "PLC Tag Name", "w": 30, "align": "L"},
    {"id": "pid", "name": "P&ID No.", "w": 25, "align": "C"},
    {"id": "desc", "name": "Service / Equipment Description", "w": 55, "align": "L"},
    {"id": "zone_flr", "name": "Zone / Floor", "w": 23, "align": "L"},
]

def render_table_header(pdf):
    pdf.set_font("Helvetica", "B", 7.5)
    pdf.set_fill_color(35, 55, 85) # Slate Navy
    pdf.set_text_color(255, 255, 255)
    
    for i, col in enumerate(COLUMNS):
        is_last = (i == len(COLUMNS) - 1)
        pdf.cell(col["w"], 6, col["name"], border=1, 
                 new_x=XPos.LMARGIN if is_last else XPos.RIGHT, 
                 new_y=YPos.NEXT if is_last else YPos.TOP, 
                 align="C", fill=True)

def truncate_text(pdf, text, max_w, font_name="Helvetica", font_style="", font_size=7):
    pdf.set_font(font_name, font_style, font_size)
    s = str(text)
    if pdf.get_string_width(s) <= max_w:
        return s
    while len(s) > 3 and pdf.get_string_width(s + "...") > max_w:
        s = s[:-1]
    return s + "..."

def render_jb_rows(pdf, df_sub, start_no=1):
    row_height = 5.2
    cur_no = start_no
    
    for _, row in df_sub.iterrows():
        if pdf.get_y() > 192:
            pdf.add_page()
            render_table_header(pdf)

        inst_tag = clean_val(row.get("Instruement Tag"))
        is_spare = inst_tag.lower() == "spare" or "spare" in inst_tag.lower()
        
        if is_spare:
            if (cur_no % 2) == 0:
                pdf.set_fill_color(248, 248, 248)
            else:
                pdf.set_fill_color(242, 242, 242)
            pdf.set_text_color(120, 120, 120)
        else:
            if (cur_no % 2) == 0:
                pdf.set_fill_color(255, 255, 255)
            else:
                pdf.set_fill_color(246, 249, 253)
            pdf.set_text_color(20, 20, 20)

        chassis = clean_val(row.get("Chassis"))
        slot = clean_val(row.get("Slot"))
        rack_slot = f"{chassis}/S{slot}" if chassis != "-" and slot != "-" else (chassis if chassis != "-" else slot)
        
        channel = clean_val(row.get("Column7"))
        terminal = clean_val(row.get("Terminal"))
        card = clean_val(row.get("Card"))
        io_type = clean_val(row.get("I/O Type"))
        signal = clean_val(row.get("Signal Type2"))
        plc_tag = clean_val(row.get("PLC_Tag"))
        pid = clean_val(row.get("P&ID No. 3.5"))
        if pid == "-":
            pid = clean_val(row.get("P&ID No. 3.3"))
            
        desc = clean_val(row.get("Instruement_Description"))
        if desc == "-" or not desc:
            desc = clean_val(row.get("Control Description"))
        if is_spare and desc == "-":
            desc = "Spare I/O Point"

        zone = clean_val(row.get("Zone"))
        flr = clean_val(row.get("Floor"))
        zone_flr = f"{zone} / {flr}" if flr != "-" else zone

        row_data = {
            "no": str(cur_no),
            "terminal": terminal,
            "slot": rack_slot,
            "channel": channel,
            "card": card,
            "type": io_type,
            "signal": signal,
            "tag": inst_tag,
            "plc_tag": plc_tag,
            "pid": pid,
            "desc": desc,
            "zone_flr": zone_flr
        }

        for i, col in enumerate(COLUMNS):
            col_id = col["id"]
            col_w = col["w"]
            align = col["align"]
            val = row_data[col_id]
            is_last = (i == len(COLUMNS) - 1)
            
            if col_id in ["tag", "plc_tag"]:
                pdf.set_font("Helvetica", "B" if not is_spare else "", 7)
            elif col_id in ["type"]:
                pdf.set_font("Helvetica", "B", 7)
            else:
                pdf.set_font("Helvetica", "", 7)
                
            fit_val = truncate_text(pdf, val, col_w - 1.5, "Helvetica", "B" if col_id in ["tag", "plc_tag", "type"] and not is_spare else "", 7)
            pdf.cell(col_w, row_height, fit_val, border=1, 
                     new_x=XPos.LMARGIN if is_last else XPos.RIGHT, 
                     new_y=YPos.NEXT if is_last else YPos.TOP, 
                     align=align, fill=True)
            
        cur_no += 1

    return cur_no

jb_stats_all = {}
destinations = sorted(df_io["Destination"].unique(), key=lambda x: str(x))

for jb in destinations:
    sub = df_io[df_io["Destination"] == jb]
    tot = len(sub)
    sp = len(sub[sub["Instruement Tag"].astype(str).str.lower().str.contains("spare")])
    act = tot - sp
    di = len(sub[sub["I/O Type"] == "DI"])
    do = len(sub[sub["I/O Type"] == "DO"])
    ai = len(sub[sub["I/O Type"] == "AI"])
    ao = len(sub[sub["I/O Type"] == "AO"])
    bus = len(sub[sub["I/O Type"] == "BUS"])
    
    jb_stats_all[jb] = {
        "Total": tot,
        "Active": act,
        "Spare": sp,
        "DI": di,
        "DO": do,
        "AI": ai,
        "AO": ao,
        "BUS": bus
    }

print("\n--- Generating Individual PDF Reports per Junction Box ---")
individual_pdf_files = []

for jb in destinations:
    sub = df_io[df_io["Destination"] == jb]
    info = jb_metadata.get(jb, {
        "location": clean_val(sub["Junction Box Location"].iloc[0] if len(sub["Junction Box Location"].dropna()) > 0 else "-"),
        "size": "-",
        "type": "Junction Box"
    })
    stats = jb_stats_all[jb]
    
    pdf = IndustrialIOPDF(jb_name=jb, jb_info=info, is_master=False, stats=stats)
    pdf.add_page()
    render_table_header(pdf)
    render_jb_rows(pdf, sub)
    
    safe_jb_name = jb.replace("/", "_").replace("\\", "_").replace(" ", "_")
    filename = f"JB_Report_{safe_jb_name}.pdf"
    filepath = os.path.join(OUTPUT_DIR, filename)
    pdf.output(filepath)
    individual_pdf_files.append((jb, filename, filepath, stats["Total"]))
    print(f"Generated: {filename:30} ({stats['Total']:3} channels)")

print("\n--- Generating Master Consolidated PDF Report ---")
master_pdf_path = os.path.join(OUTPUT_DIR, "IO_List_All_Junction_Boxes_Master_Report.pdf")
master_pdf = IndustrialIOPDF(jb_name="Master Report", jb_info={}, is_master=True, stats={})

# --- Cover Page ---
master_pdf.add_page()

# Decorative top band
master_pdf.set_fill_color(24, 43, 73)
master_pdf.rect(0, 0, 297, 28, "F")

master_pdf.set_text_color(255, 255, 255)
master_pdf.set_font("Helvetica", "B", 16)
master_pdf.set_xy(10, 8)
master_pdf.cell(0, 10, "INGREDION (THAILAND) CO., LTD.  |  KALASIN PLANT", 0, new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="C")

master_pdf.ln(25)
master_pdf.set_text_color(24, 43, 73)
master_pdf.set_font("Helvetica", "B", 24)
master_pdf.cell(0, 12, "PROJECT JET COOKER", 0, new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="C")

master_pdf.set_font("Helvetica", "B", 18)
master_pdf.set_text_color(70, 80, 95)
master_pdf.cell(0, 10, "INSTRUMENT I/O LIST - JUNCTION BOX MASTER REPORT", 0, new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="C")

master_pdf.ln(6)
master_pdf.set_draw_color(200, 210, 225)
master_pdf.set_line_width(0.8)
master_pdf.line(40, master_pdf.get_y(), 257, master_pdf.get_y())
master_pdf.ln(8)

# Document Details Box
master_pdf.set_fill_color(246, 249, 253)
master_pdf.set_draw_color(180, 200, 220)
master_pdf.set_line_width(0.3)
box_x, box_y, box_w, box_h = 50, master_pdf.get_y(), 197, 45
master_pdf.rect(box_x, box_y, box_w, box_h, "DF")

master_pdf.set_xy(box_x + 10, box_y + 6)
master_pdf.set_font("Helvetica", "B", 10)
master_pdf.set_text_color(30, 40, 60)
master_pdf.cell(80, 7, "Document Number:", 0, new_x=XPos.RIGHT, new_y=YPos.TOP)
master_pdf.set_font("Helvetica", "", 10)
master_pdf.cell(90, 7, "Instrument I/O List (Kalasin Rev. 3.6)", 0, new_x=XPos.LMARGIN, new_y=YPos.NEXT)

master_pdf.set_x(box_x + 10)
master_pdf.set_font("Helvetica", "B", 10)
master_pdf.cell(80, 7, "Engineering Contractor:", 0, new_x=XPos.RIGHT, new_y=YPos.TOP)
master_pdf.set_font("Helvetica", "", 10)
master_pdf.cell(90, 7, "AEC Industrial Engineering Co., Ltd.", 0, new_x=XPos.LMARGIN, new_y=YPos.NEXT)

master_pdf.set_x(box_x + 10)
master_pdf.set_font("Helvetica", "B", 10)
master_pdf.cell(80, 7, "Total Enclosures Covered:", 0, new_x=XPos.RIGHT, new_y=YPos.TOP)
master_pdf.set_font("Helvetica", "", 10)
master_pdf.cell(90, 7, f"{len(destinations)} Enclosures (14 Field JBs, CA1 Cabinet, MCC)", 0, new_x=XPos.LMARGIN, new_y=YPos.NEXT)

master_pdf.set_x(box_x + 10)
master_pdf.set_font("Helvetica", "B", 10)
master_pdf.cell(80, 7, "Total Configured I/O Points:", 0, new_x=XPos.RIGHT, new_y=YPos.TOP)
master_pdf.set_font("Helvetica", "", 10)
tot_active = sum(s["Active"] for s in jb_stats_all.values())
tot_spare = sum(s["Spare"] for s in jb_stats_all.values())
tot_all = sum(s["Total"] for s in jb_stats_all.values())
master_pdf.cell(90, 7, f"{tot_all} Points ({tot_active} Active, {tot_spare} Spares - {tot_spare/tot_all*100:.1f}%)", 0, new_x=XPos.LMARGIN, new_y=YPos.NEXT)

master_pdf.set_x(box_x + 10)
master_pdf.set_font("Helvetica", "B", 10)
master_pdf.cell(80, 7, "Report Issue Date:", 0, new_x=XPos.RIGHT, new_y=YPos.TOP)
master_pdf.set_font("Helvetica", "", 10)
master_pdf.cell(90, 7, "2026-09-07", 0, new_x=XPos.LMARGIN, new_y=YPos.NEXT)

master_pdf.set_y(box_y + box_h + 15)
master_pdf.set_font("Helvetica", "I", 9)
master_pdf.set_text_color(110, 120, 130)
master_pdf.cell(0, 6, "This document compiles the detailed signal wiring schedule and terminal assignments for all junction boxes.", 0, new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="C")

# --- Executive Summary Schedule Table ---
master_pdf.is_summary_page = True
master_pdf.add_page()
master_pdf.set_font("Helvetica", "B", 13)
master_pdf.set_text_color(24, 43, 73)
master_pdf.cell(0, 8, "JUNCTION BOX & ENCLOSURE EXECUTIVE SUMMARY SCHEDULE", 0, new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="L")
master_pdf.ln(2)

SUM_COLS = [
    {"name": "No.", "w": 10, "align": "C"},
    {"name": "Enclosure Tag", "w": 25, "align": "C"},
    {"name": "Type / Classification", "w": 52, "align": "L"},
    {"name": "Location", "w": 45, "align": "L"},
    {"name": "Dimension (mm)", "w": 25, "align": "C"},
    {"name": "Total", "w": 15, "align": "C"},
    {"name": "Active", "w": 15, "align": "C"},
    {"name": "Spare", "w": 15, "align": "C"},
    {"name": "DI", "w": 15, "align": "C"},
    {"name": "DO", "w": 15, "align": "C"},
    {"name": "AI", "w": 15, "align": "C"},
    {"name": "AO", "w": 15, "align": "C"},
    {"name": "BUS", "w": 15, "align": "C"},
]

master_pdf.set_font("Helvetica", "B", 7.5)
master_pdf.set_fill_color(35, 55, 85)
master_pdf.set_text_color(255, 255, 255)
for i, col in enumerate(SUM_COLS):
    is_last = (i == len(SUM_COLS) - 1)
    master_pdf.cell(col["w"], 6.5, col["name"], border=1, 
                    new_x=XPos.LMARGIN if is_last else XPos.RIGHT, 
                    new_y=YPos.NEXT if is_last else YPos.TOP, 
                    align="C", fill=True)

s_no = 1
for jb in destinations:
    info = jb_metadata.get(jb, {})
    st = jb_stats_all[jb]
    loc = info.get("location", "-")
    size = info.get("size", "-")
    jtype = info.get("type", "Standard JB")
    
    if s_no % 2 == 0:
        master_pdf.set_fill_color(248, 250, 252)
    else:
        master_pdf.set_fill_color(255, 255, 255)
        
    master_pdf.set_text_color(30, 30, 30)
    master_pdf.set_font("Helvetica", "", 7.5)
    
    master_pdf.cell(10, 5.8, str(s_no), border=1, new_x=XPos.RIGHT, new_y=YPos.TOP, align="C", fill=True)
    master_pdf.set_font("Helvetica", "B", 7.5)
    master_pdf.cell(25, 5.8, jb, border=1, new_x=XPos.RIGHT, new_y=YPos.TOP, align="C", fill=True)
    master_pdf.set_font("Helvetica", "", 7.5)
    master_pdf.cell(52, 5.8, truncate_text(master_pdf, jtype, 50, "Helvetica", "", 7.5), border=1, new_x=XPos.RIGHT, new_y=YPos.TOP, align="L", fill=True)
    master_pdf.cell(45, 5.8, truncate_text(master_pdf, loc, 43, "Helvetica", "", 7.5), border=1, new_x=XPos.RIGHT, new_y=YPos.TOP, align="L", fill=True)
    master_pdf.cell(25, 5.8, size, border=1, new_x=XPos.RIGHT, new_y=YPos.TOP, align="C", fill=True)
    
    master_pdf.set_font("Helvetica", "B", 7.5)
    master_pdf.cell(15, 5.8, str(st["Total"]), border=1, new_x=XPos.RIGHT, new_y=YPos.TOP, align="C", fill=True)
    master_pdf.set_font("Helvetica", "", 7.5)
    master_pdf.cell(15, 5.8, str(st["Active"]), border=1, new_x=XPos.RIGHT, new_y=YPos.TOP, align="C", fill=True)
    master_pdf.cell(15, 5.8, str(st["Spare"]), border=1, new_x=XPos.RIGHT, new_y=YPos.TOP, align="C", fill=True)
    master_pdf.cell(15, 5.8, str(st["DI"]), border=1, new_x=XPos.RIGHT, new_y=YPos.TOP, align="C", fill=True)
    master_pdf.cell(15, 5.8, str(st["DO"]), border=1, new_x=XPos.RIGHT, new_y=YPos.TOP, align="C", fill=True)
    master_pdf.cell(15, 5.8, str(st["AI"]), border=1, new_x=XPos.RIGHT, new_y=YPos.TOP, align="C", fill=True)
    master_pdf.cell(15, 5.8, str(st["AO"]), border=1, new_x=XPos.RIGHT, new_y=YPos.TOP, align="C", fill=True)
    master_pdf.cell(15, 5.8, str(st["BUS"]), border=1, new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="C", fill=True)
    s_no += 1

# Total row
master_pdf.set_fill_color(225, 235, 245)
master_pdf.set_text_color(20, 30, 50)
master_pdf.set_font("Helvetica", "B", 8)
master_pdf.cell(157, 6.5, "TOTAL (ALL ENCLOSURES)", border=1, new_x=XPos.RIGHT, new_y=YPos.TOP, align="C", fill=True)
master_pdf.cell(15, 6.5, str(tot_all), border=1, new_x=XPos.RIGHT, new_y=YPos.TOP, align="C", fill=True)
master_pdf.cell(15, 6.5, str(tot_active), border=1, new_x=XPos.RIGHT, new_y=YPos.TOP, align="C", fill=True)
master_pdf.cell(15, 6.5, str(tot_spare), border=1, new_x=XPos.RIGHT, new_y=YPos.TOP, align="C", fill=True)
master_pdf.cell(15, 6.5, str(sum(s["DI"] for s in jb_stats_all.values())), border=1, new_x=XPos.RIGHT, new_y=YPos.TOP, align="C", fill=True)
master_pdf.cell(15, 6.5, str(sum(s["DO"] for s in jb_stats_all.values())), border=1, new_x=XPos.RIGHT, new_y=YPos.TOP, align="C", fill=True)
master_pdf.cell(15, 6.5, str(sum(s["AI"] for s in jb_stats_all.values())), border=1, new_x=XPos.RIGHT, new_y=YPos.TOP, align="C", fill=True)
master_pdf.cell(15, 6.5, str(sum(s["AO"] for s in jb_stats_all.values())), border=1, new_x=XPos.RIGHT, new_y=YPos.TOP, align="C", fill=True)
master_pdf.cell(15, 6.5, str(sum(s["BUS"] for s in jb_stats_all.values())), border=1, new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="C", fill=True)
master_pdf.ln(10)

# --- Detailed JB Sections in Master Report ---
master_pdf.is_summary_page = False
for jb in destinations:
    sub = df_io[df_io["Destination"] == jb]
    info = jb_metadata.get(jb, {})
    stats = jb_stats_all[jb]
    
    master_pdf.jb_name = jb
    master_pdf.jb_info = info
    master_pdf.stats = stats
    
    master_pdf.add_page()
    render_table_header(master_pdf)
    render_jb_rows(master_pdf, sub)

master_pdf.output(master_pdf_path)
print(f"\nSuccessfully generated Master Report: {master_pdf_path}")
print(f"Total pages in Master Report: {master_pdf.page_no()}")
print("ALL PDF REPORTS GENERATED SUCCESSFULLY!")
