#!/usr/bin/env python3
"""
================================================================================
Kalasin Engineering xCIP System (Project Ref: x2608003)
CHASSIS C1-C8 SLOT TO DESTINATION JUNCTION BOX MASTER REPORT
================================================================================
Source: 
  - 03_IO_Lists_and_Schedules/IO_List_xDev-R01-Tag35-7.xlsx (Sheet: 'IO List')
Outputs:
  - 03_IO_Lists_and_Schedules/C1_C8_Chassis_Slot_to_Destination_JB_Report.xlsx

Styling: Light theme, engineering executive standard, openpyxl professional formatting.
================================================================================
"""

import os
import sys
from collections import defaultdict, Counter
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = "/Users/x92120/xApp-001/x260911-002-KalasinEngineering"
SOURCE_EXCEL = os.path.join(BASE_DIR, "03_IO_Lists_and_Schedules/IO_List_xDev-R01-Tag35-7.xlsx")
OUTPUT_EXCEL = os.path.join(BASE_DIR, "03_IO_Lists_and_Schedules/C1_C8_Chassis_Slot_to_Destination_JB_Report.xlsx")

# Theme Palette (Light Executive Standard)
NAVY_HEADER = "1E293B"      # Deep Slate Navy
NAVY_SUBHEADER = "334155"   # Slate Subheader
WHITE = "FFFFFF"
LIGHT_BG_1 = "FFFFFF"
LIGHT_BG_2 = "F8FAFC"       # Very light slate
BORDER_COLOR = "CBD5E1"     # Light slate border

# Group Header Colors
SRC_HDR_FILL = "DBEAFE"     # Light Ice Blue
SRC_HDR_FONT = "1E3A8A"     # Deep Blue
DEST_HDR_FILL = "FEF3C7"    # Light Amber/Sand
DEST_HDR_FONT = "78350F"    # Deep Amber/Brown
SYS_HDR_FILL = "E2E8F0"     # Pale Slate
SYS_HDR_FONT = "0F172A"     # Slate 900

# Status Fills
ACTIVE_FILL = "DCFCE7"      # Light Emerald
ACTIVE_FONT = "14532D"      # Dark Forest Green
SPARE_FILL = "FEF9C3"       # Light Yellow/Amber
SPARE_FONT = "854D0E"       # Dark Amber

CHASSIS_INFO = {
    'C1': {
        'name': 'Main Controller Rack (Chassis 1)',
        'panel': 'Control Cabinet CA1 (Main Control Room)',
        'role': '1756-L950TPSXT Controller + DLR Supervisors + Local Fast I/O',
        'network': 'Stratix 5400 / EN4TR Ring Node 1 (192.168.1.10)',
        'slots_total': 13
    },
    'C2': {
        'name': 'Main Expansion Rack (Chassis 2)',
        'panel': 'Control Cabinet CA1 (Main Control Room)',
        'role': 'High-Density DI/DO/AI Expansion for Spray Dryer & Jet Cooker',
        'network': '1756-EN4TR Ring Node 2 (192.168.1.11)',
        'slots_total': 13
    },
    'C3': {
        'name': 'Remote I/O Rack 1 (Chassis 3)',
        'panel': 'RIO Cabinet / Field Enclosure (Spray Dryer 3rd/7th Floor)',
        'role': 'Field I/O (DI, DO, AI, RTD/AO) for Packing Tower & Spray Dryer',
        'network': '1756-EN4TR DLR Field Drop',
        'slots_total': 13
    },
    'C4': {
        'name': 'Remote I/O Rack 2 (Chassis 4)',
        'panel': 'RIO Cabinet / Field Enclosure (Spray Dryer 6th/8th Floor)',
        'role': 'Field I/O (DI, DO, AI) for Upper Process Levels & Packing Tower',
        'network': '1756-EN4TR DLR Field Drop',
        'slots_total': 13
    },
    'C5': {
        'name': 'Remote I/O Station 5 - RIO-200 (Chassis 5)',
        'panel': 'Slurry Building Enclosure (Slurry - Out Building 2nd Floor)',
        'role': 'Dedicated Out-Building Slurry Process Control (DI, DO, AI)',
        'network': '1756-EN4TR Fiber Optic Drop',
        'slots_total': 7
    },
    'C6': {
        'name': 'Motor Control Interface (Chassis 6)',
        'panel': 'Motor Control Center (MCC Room Ground Floor)',
        'role': 'Direct Hardwired Interlocks & Permissives for Heavy Drives',
        'network': 'Dedicated Sub-Rack / Communication Link',
        'slots_total': 7
    },
    'C7': {
        'name': 'MCC & Bus Supervisory Rack (Chassis 7)',
        'panel': 'Motor Control Center (MCC Room Ground Floor)',
        'role': 'Fieldbus Gateway, VFD Drives Supervision & High-Density MCC Control',
        'network': 'EtherNet/IP to PowerFlex VFDs / Bus Link',
        'slots_total': 13
    },
    'C8': {
        'name': 'Intrinsically Safe Sub-System (Chassis 8)',
        'panel': 'IS Galvanic Barrier Enclosures (Field Ex Zones)',
        'role': 'ATEX/Ex Zone Hazard Boundary Instruments (IS-JB-603/608/612/618)',
        'network': 'IS Isolated Signal Conditioning Racks',
        'slots_total': 10
    }
}

JB_INFO = {
    'CA1': {'loc': 'Main Control Room', 'desc': 'Main Control Automation Cabinet (Central PLC Cabinets)'},
    'JB-401': {'loc': 'Infeed 2nd Floor', 'desc': 'Infeed Process Junction Box (Raw Materials & Wet Prep)'},
    'JB-402': {'loc': 'Jet Cooker 2nd Floor', 'desc': 'Jet Cooker & Thermal Cooking Junction Box'},
    'JB-601': {'loc': 'Spray Dryer 1st Floor', 'desc': 'Spray Dryer Base & Atomizer Discharge Junction Box'},
    'JB-602': {'loc': 'Spray Dryer 3rd Floor', 'desc': 'Spray Dryer Chamber Lower Junction Box'},
    'JB-606': {'loc': 'Spray Dryer 6th Floor', 'desc': 'Spray Dryer Middle Body & Filter Junction Box'},
    'JB-607': {'loc': 'Spray Dryer 7th Floor', 'desc': 'Spray Dryer Cyclone & Exhaust System Junction Box'},
    'JB-608': {'loc': 'Spray Dryer 8th Floor', 'desc': 'Spray Dryer Top Air Dispenser & Scrubber Junction Box'},
    'JB-612': {'loc': 'Packing Tower 2nd Floor', 'desc': 'Packing Tower Lower Transfer Junction Box'},
    'JB-618': {'loc': 'Packing Tower 8th Floor', 'desc': 'Packing Tower Upper Cyclone & Dehumidifier Junction Box'},
    'RIO-200': {'loc': 'Slurry Out-Building 2nd Fl', 'desc': 'Remote Slurry Processing Building RIO Station'},
    'MCC': {'loc': 'MCC Room Ground Floor', 'desc': 'Main Motor Control Center & VFD Inverter Switchgear'},
    'IS-JB-603': {'loc': 'Spray Dryer 3rd Floor (IS)', 'desc': 'Intrinsically Safe Junction Box (Ex Rated Chamber)'},
    'IS-JB-608': {'loc': 'Spray Dryer 8th Floor (IS)', 'desc': 'Intrinsically Safe Junction Box (Ex Top Level)'},
    'IS-JB-612': {'loc': 'Packing Tower 2nd Floor (IS)', 'desc': 'Intrinsically Safe Junction Box (Ex Lower Tower)'},
    'IS-JB-618': {'loc': 'Packing Tower 8th Floor (IS)', 'desc': 'Intrinsically Safe Junction Box (Ex Upper Tower)'}
}

def thin_border():
    s = Side(style='thin', color=BORDER_COLOR)
    return Border(left=s, right=s, top=s, bottom=s)

def header_border():
    s_light = Side(style='thin', color="475569")
    s_bottom = Side(style='medium', color="0F172A")
    return Border(left=s_light, right=s_light, top=s_light, bottom=s_bottom)

def load_data():
    wb = openpyxl.load_workbook(SOURCE_EXCEL, data_only=True)
    ws = wb['IO List']
    
    records = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        col7 = str(r[2]).strip() if r[2] is not None else ''
        ch = str(r[30]).strip() if r[30] is not None else ''
        slot = str(r[31]).strip() if r[31] is not None else ''
        if not ch and col7.startswith('C') and 'S' in col7:
            ch = col7.split('S')[0]
            slot = col7.split('S')[1].split(',')[0]
            
        if ch not in [f'C{i}' for i in range(1, 9)]:
            continue
            
        dest = str(r[1]).strip() if r[1] is not None else 'UNSPECIFIED'
        card = str(r[3]).strip() if r[3] is not None else ''
        card_type = str(r[26]).strip() if r[26] is not None else ''
        terminal = str(r[5]).strip() if r[5] is not None else ''
        term_label = str(r[7]).strip() if r[7] is not None else ''
        jb_loc = str(r[8]).strip() if r[8] is not None else ''
        tag = str(r[11]).strip() if r[11] is not None else ''
        desc = str(r[12]).strip() if r[12] is not None else ''
        floor = str(r[13]).strip() if r[13] is not None else ''
        pid = str(r[14]).strip() if r[14] is not None else ''
        plc_tag = str(r[19]).strip() if r[19] is not None else ''
        io_type = str(r[20]).strip() if r[20] is not None else ''
        point = str(r[32]).strip() if r[32] is not None else ''
        
        # Clean Card Model
        model = card if card and card != '#N/A' else card_type
        if model and model != '#N/A':
            # normalize 1756-IB32-0 to 1756-IB32
            parts = model.split('-')
            if len(parts) >= 2 and parts[1].isdigit() is False:
                clean_model = f"{parts[0]}-{parts[1]}"
            else:
                clean_model = model
        else:
            # Fallback by io_type
            if io_type == 'DI': clean_model = '1756-IB32'
            elif io_type == 'DO': clean_model = '1756-OB32'
            elif io_type == 'AI': clean_model = '1756-IF16'
            elif io_type == 'AO': clean_model = '1756-OF8'
            elif io_type == 'BUS': clean_model = 'Fieldbus Comm'
            elif io_type == 'CPU': clean_model = '1756-L950TPSXT'
            elif 'ETH' in io_type: clean_model = '1756-EN4TR'
            else: clean_model = 'Module-Std'
            
        is_spare = 'SPARE' in tag.upper() or not tag or tag == '#N/A' or tag == '0'
        
        rec = {
            'chassis': ch,
            'slot': int(slot) if slot.isdigit() else slot,
            'slot_str': f"Slot {int(slot):02d}" if str(slot).isdigit() else str(slot),
            'point': int(point) if str(point).isdigit() else point,
            'col7': col7,
            'dest': dest,
            'card_model': clean_model,
            'terminal': terminal,
            'term_label': term_label,
            'jb_loc': jb_loc,
            'tag': tag if not is_spare else 'SPARE',
            'desc': desc if (desc and desc != '#N/A' and desc != '0') else ('Spare Reserve Channel' if is_spare else '-'),
            'floor': floor if (floor and floor != '#N/A') else '-',
            'pid': pid if (pid and pid != '#N/A') else '-',
            'plc_tag': plc_tag if (plc_tag and plc_tag != '#N/A') else '-',
            'io_type': io_type,
            'is_spare': is_spare
        }
        records.append(rec)
    return records

def build_report():
    print(f"Loading data from: {SOURCE_EXCEL}")
    records = load_data()
    print(f"Total valid C1-C8 I/O records: {len(records)}")
    
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)
    
    # -------------------------------------------------------------
    # Sheet 1: Executive Dashboard & Summary
    # -------------------------------------------------------------
    ws_dash = wb.create_sheet(title="00_Executive_Summary")
    ws_dash.views.sheetView[0].showGridLines = True
    
    # Title Block
    ws_dash.merge_cells("A1:K1")
    tcell = ws_dash["A1"]
    tcell.value = "KALASIN ENGINEERING xCIP-1545 PROJECT (Ref: x2608003)"
    tcell.font = Font(name="Arial", size=14, bold=True, color=WHITE)
    tcell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    tcell.alignment = Alignment(horizontal="center", vertical="center")
    ws_dash.row_dimensions[1].height = 28
    
    ws_dash.merge_cells("A2:K2")
    subcell = ws_dash["A2"]
    subcell.value = "CHASSIS C1 - C8 TO DESTINATION JUNCTION BOX ROUTING DIRECTORY"
    subcell.font = Font(name="Arial", size=11, bold=True, color=WHITE)
    subcell.fill = PatternFill(start_color=NAVY_SUBHEADER, end_color=NAVY_SUBHEADER, fill_type="solid")
    subcell.alignment = Alignment(horizontal="center", vertical="center")
    ws_dash.row_dimensions[2].height = 22
    
    ws_dash.append([])
    
    # Section 1: Chassis Summary Table
    ws_dash.append(["1. CHASSIS SYSTEM SUMMARY (C1 - C8 OVERVIEW)"])
    ws_dash.merge_cells("A4:K4")
    c4 = ws_dash["A4"]
    c4.font = Font(name="Arial", size=11, bold=True, color=SRC_HDR_FONT)
    c4.fill = PatternFill(start_color=SRC_HDR_FILL, end_color=SRC_HDR_FILL, fill_type="solid")
    c4.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_dash.row_dimensions[4].height = 24
    
    headers_c = [
        "Chassis ID", "Chassis Description", "Physical Location / Enclosure",
        "Configured Slots", "Total Points", "Active Points", "Spare Points", 
        "DI", "DO", "AI", "AO / BUS"
    ]
    ws_dash.append(headers_c)
    ws_dash.row_dimensions[5].height = 22
    for col_idx, h in enumerate(headers_c, start=1):
        cell = ws_dash.cell(row=5, column=col_idx)
        cell.font = Font(name="Arial", size=9, bold=True, color=WHITE)
        cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = header_border()
        
    chassis_stats = defaultdict(lambda: {
        'total': 0, 'active': 0, 'spare': 0, 
        'DI': 0, 'DO': 0, 'AI': 0, 'AO': 0, 'BUS': 0,
        'slots': set(), 'dests': set()
    })
    for r in records:
        ch = r['chassis']
        chassis_stats[ch]['total'] += 1
        if r['is_spare']:
            chassis_stats[ch]['spare'] += 1
        else:
            chassis_stats[ch]['active'] += 1
        io = r['io_type']
        if io in chassis_stats[ch]:
            chassis_stats[ch][io] += 1
        chassis_stats[ch]['slots'].add(r['slot'])
        chassis_stats[ch]['dests'].add(r['dest'])
        
    row_num = 6
    for ch_idx in range(1, 9):
        ch = f"C{ch_idx}"
        st = chassis_stats[ch]
        info = CHASSIS_INFO.get(ch, {})
        row_vals = [
            ch,
            info.get('name', f"Chassis {ch}"),
            info.get('panel', 'Field Rack'),
            len(st['slots']),
            st['total'],
            st['active'],
            st['spare'],
            st['DI'],
            st['DO'],
            st['AI'],
            st['AO'] + st['BUS']
        ]
        ws_dash.append(row_vals)
        ws_dash.row_dimensions[row_num].height = 20
        fill_col = LIGHT_BG_1 if row_num % 2 == 0 else LIGHT_BG_2
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws_dash.cell(row=row_num, column=col_idx)
            cell.font = Font(name="Arial", size=9, bold=(col_idx == 1))
            cell.fill = PatternFill(start_color=fill_col, end_color=fill_col, fill_type="solid")
            cell.border = thin_border()
            if col_idx in [1, 4]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [2, 3]:
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")
        row_num += 1
        
    # Total Row
    tot_row_idx = row_num
    ws_dash.append([
        "TOTAL", "All Chassis (C1 - C8 Combined)", "-", 
        sum(len(chassis_stats[f'C{i}']['slots']) for i in range(1, 9)),
        len(records),
        sum(chassis_stats[f'C{i}']['active'] for i in range(1, 9)),
        sum(chassis_stats[f'C{i}']['spare'] for i in range(1, 9)),
        sum(chassis_stats[f'C{i}']['DI'] for i in range(1, 9)),
        sum(chassis_stats[f'C{i}']['DO'] for i in range(1, 9)),
        sum(chassis_stats[f'C{i}']['AI'] for i in range(1, 9)),
        sum(chassis_stats[f'C{i}']['AO'] + chassis_stats[f'C{i}']['BUS'] for i in range(1, 9))
    ])
    ws_dash.row_dimensions[tot_row_idx].height = 22
    for col_idx in range(1, 12):
        cell = ws_dash.cell(row=tot_row_idx, column=col_idx)
        cell.font = Font(name="Arial", size=9, bold=True, color=WHITE)
        cell.fill = PatternFill(start_color=NAVY_SUBHEADER, end_color=NAVY_SUBHEADER, fill_type="solid")
        cell.border = header_border()
        if col_idx in [1, 4]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_idx in [2, 3]:
            cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        else:
            cell.alignment = Alignment(horizontal="right", vertical="center")
            
    # Section 2: Destination Junction Box Directory
    ws_dash.append([])
    sec2_row = tot_row_idx + 2
    ws_dash.append(["2. DESTINATION JUNCTION BOX & CABINET DIRECTORY"])
    ws_dash.merge_cells(f"A{sec2_row}:K{sec2_row}")
    c_sec2 = ws_dash[f"A{sec2_row}"]
    c_sec2.font = Font(name="Arial", size=11, bold=True, color=DEST_HDR_FONT)
    c_sec2.fill = PatternFill(start_color=DEST_HDR_FILL, end_color=DEST_HDR_FILL, fill_type="solid")
    c_sec2.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_dash.row_dimensions[sec2_row].height = 24
    
    headers_d = [
        "Destination JB / Cabinet", "Plant Location", "Description / Process Role",
        "Feeding Chassis", "Connected Slots", "Total Signals", "Active Signals", 
        "Spare Signals", "DI", "DO", "AI/AO/BUS"
    ]
    ws_dash.append(headers_d)
    sec2_hdr_row = sec2_row + 1
    ws_dash.row_dimensions[sec2_hdr_row].height = 22
    for col_idx, h in enumerate(headers_d, start=1):
        cell = ws_dash.cell(row=sec2_hdr_row, column=col_idx)
        cell.font = Font(name="Arial", size=9, bold=True, color=WHITE)
        cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = header_border()
        
    jb_stats = defaultdict(lambda: {
        'total': 0, 'active': 0, 'spare': 0,
        'DI': 0, 'DO': 0, 'AI': 0, 'AO': 0, 'BUS': 0,
        'chassis': set(), 'slots': set()
    })
    for r in records:
        d = r['dest']
        jb_stats[d]['total'] += 1
        if r['is_spare']:
            jb_stats[d]['spare'] += 1
        else:
            jb_stats[d]['active'] += 1
        io = r['io_type']
        if io in jb_stats[d]:
            jb_stats[d][io] += 1
        jb_stats[d]['chassis'].add(r['chassis'])
        jb_stats[d]['slots'].add(f"{r['chassis']}S{r['slot']:02d}")
        
    d_row = sec2_hdr_row + 1
    for d in sorted(jb_stats.keys()):
        st = jb_stats[d]
        info = JB_INFO.get(d, {})
        ch_str = ', '.join(sorted(st['chassis'], key=lambda x: int(x[1:])))
        slots_str = ', '.join(sorted(st['slots']))
        if len(slots_str) > 40:
            slots_str = f"{len(st['slots'])} slots ({slots_str[:35]}...)"
            
        row_vals = [
            d,
            info.get('loc', 'Plant Field Area'),
            info.get('desc', 'Process Field Junction Box'),
            ch_str,
            slots_str,
            st['total'],
            st['active'],
            st['spare'],
            st['DI'],
            st['DO'],
            st['AI'] + st['AO'] + st['BUS']
        ]
        ws_dash.append(row_vals)
        ws_dash.row_dimensions[d_row].height = 20
        fill_col = LIGHT_BG_1 if d_row % 2 == 0 else LIGHT_BG_2
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws_dash.cell(row=d_row, column=col_idx)
            cell.font = Font(name="Arial", size=9, bold=(col_idx == 1))
            cell.fill = PatternFill(start_color=fill_col, end_color=fill_col, fill_type="solid")
            cell.border = thin_border()
            if col_idx in [1, 4]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [2, 3, 5]:
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")
        d_row += 1

    # -------------------------------------------------------------
    # Sheet 2: C1-C8 Slot-to-Destination Schedule (Aggregated by Slot & JB)
    # -------------------------------------------------------------
    ws_sched = wb.create_sheet(title="01_Slot_to_JB_Schedule")
    ws_sched.views.sheetView[0].showGridLines = True
    
    ws_sched.merge_cells("A1:K1")
    t2 = ws_sched["A1"]
    t2.value = "CHASSIS C1 TO C8: COMPLETE SLOT-TO-DESTINATION ROUTING SCHEDULE"
    t2.font = Font(name="Arial", size=13, bold=True, color=WHITE)
    t2.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    t2.alignment = Alignment(horizontal="center", vertical="center")
    ws_sched.row_dimensions[1].height = 26
    
    headers_s2 = [
        "Chassis", "Slot", "Slot Tag", "Hardware Module", "Signal Type", 
        "Destination JB", "JB Location", "Total Pts", "Active Pts", "Spare Pts", "Sample Active Loops / Tags"
    ]
    ws_sched.append(headers_s2)
    ws_sched.row_dimensions[2].height = 22
    for col_idx, h in enumerate(headers_s2, start=1):
        cell = ws_sched.cell(row=2, column=col_idx)
        cell.font = Font(name="Arial", size=9, bold=True, color=WHITE)
        cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = header_border()
        
    # Group records by (chassis, slot, dest)
    csd_groups = defaultdict(list)
    for r in records:
        csd_groups[(r['chassis'], r['slot'], r['dest'])].append(r)
        
    sorted_keys = sorted(csd_groups.keys(), key=lambda x: (int(x[0][1:]), int(x[1]) if str(x[1]).isdigit() else 99, x[2]))
    
    current_ch = None
    r_idx = 3
    for ch, slot, dest in sorted_keys:
        items = csd_groups[(ch, slot, dest)]
        card = items[0]['card_model']
        io_type = items[0]['io_type']
        jb_loc = JB_INFO.get(dest, {}).get('loc', items[0]['jb_loc'])
        
        tot = len(items)
        spr = sum(1 for it in items if it['is_spare'])
        act = tot - spr
        
        active_tags = [it['tag'] for it in items if not it['is_spare']]
        sample_str = ', '.join(active_tags[:4])
        if len(active_tags) > 4:
            sample_str += f" (+{len(active_tags) - 4} more)"
        elif not sample_str:
            sample_str = "All Spares / Reserve"
            
        slot_str = f"Slot {int(slot):02d}" if str(slot).isdigit() else str(slot)
        slot_tag = f"{ch}S{int(slot):02d}" if str(slot).isdigit() else f"{ch}S{slot}"
        
        row_vals = [
            ch, slot_str, slot_tag, card, io_type, dest, jb_loc, tot, act, spr, sample_str
        ]
        ws_sched.append(row_vals)
        ws_sched.row_dimensions[r_idx].height = 20
        fill_col = LIGHT_BG_1 if r_idx % 2 == 0 else LIGHT_BG_2
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws_sched.cell(row=r_idx, column=col_idx)
            cell.font = Font(name="Arial", size=9, bold=(col_idx in [1, 2, 6]))
            cell.fill = PatternFill(start_color=fill_col, end_color=fill_col, fill_type="solid")
            cell.border = thin_border()
            if col_idx in [1, 2, 3, 5, 6]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [4, 7, 11]:
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")
        r_idx += 1

    # -------------------------------------------------------------
    # Sheet 3: Destination to Chassis/Slot Cross-Reference Matrix
    # -------------------------------------------------------------
    ws_mat = wb.create_sheet(title="02_JB_to_Chassis_Matrix")
    ws_mat.views.sheetView[0].showGridLines = True
    
    ws_mat.merge_cells("A1:K1")
    t3 = ws_mat["A1"]
    t3.value = "DESTINATION JUNCTION BOX TO CHASSIS FEEDING MATRIX"
    t3.font = Font(name="Arial", size=13, bold=True, color=WHITE)
    t3.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    t3.alignment = Alignment(horizontal="center", vertical="center")
    ws_mat.row_dimensions[1].height = 26
    
    mat_headers = ["Destination JB", "Location / Area", "C1", "C2", "C3", "C4", "C5", "C6", "C7", "C8", "Total Points"]
    ws_mat.append(mat_headers)
    ws_mat.row_dimensions[2].height = 22
    for col_idx, h in enumerate(mat_headers, start=1):
        cell = ws_mat.cell(row=2, column=col_idx)
        cell.font = Font(name="Arial", size=9, bold=True, color=WHITE)
        cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border()
        
    dest_ch_matrix = defaultdict(lambda: Counter())
    for r in records:
        dest_ch_matrix[r['dest']][r['chassis']] += 1
        
    m_row = 3
    for dest in sorted(dest_ch_matrix.keys()):
        loc = JB_INFO.get(dest, {}).get('loc', 'Field Area')
        c_counts = [dest_ch_matrix[dest][f"C{i}"] for i in range(1, 9)]
        tot_pts = sum(c_counts)
        row_vals = [dest, loc] + [cnt if cnt > 0 else "-" for cnt in c_counts] + [tot_pts]
        ws_mat.append(row_vals)
        ws_mat.row_dimensions[m_row].height = 20
        fill_col = LIGHT_BG_1 if m_row % 2 == 0 else LIGHT_BG_2
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws_mat.cell(row=m_row, column=col_idx)
            cell.font = Font(name="Arial", size=9, bold=(col_idx in [1, 11] or (col_idx >= 3 and val != "-")))
            cell.fill = PatternFill(start_color=fill_col, end_color=fill_col, fill_type="solid")
            cell.border = thin_border()
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 2:
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            else:
                cell.alignment = Alignment(horizontal="center" if val == "-" else "right", vertical="center")
        m_row += 1

    # -------------------------------------------------------------
    # Sheet 4: Complete Point-by-Point Schedule (1,218 Points)
    # -------------------------------------------------------------
    ws_full = wb.create_sheet(title="03_Master_Point_Schedule")
    ws_full.views.sheetView[0].showGridLines = True
    
    ws_full.merge_cells("A1:M1")
    t4 = ws_full["A1"]
    t4.value = "C1-C8 MASTER POINT SCHEDULE (ALL 1,218 CHANNELS WITH DESTINATION JBs)"
    t4.font = Font(name="Arial", size=13, bold=True, color=WHITE)
    t4.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    t4.alignment = Alignment(horizontal="center", vertical="center")
    ws_full.row_dimensions[1].height = 26
    
    headers_full = [
        "Chassis", "Slot", "Point", "Terminal", "Hardware Card", "Signal Type",
        "Instrument Tag", "Instrument Description", "Destination JB", "Location", 
        "Floor", "P&ID No.", "Status"
    ]
    ws_full.append(headers_full)
    ws_full.row_dimensions[2].height = 22
    for col_idx, h in enumerate(headers_full, start=1):
        cell = ws_full.cell(row=2, column=col_idx)
        cell.font = Font(name="Arial", size=9, bold=True, color=WHITE)
        cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border()
        
    f_row = 3
    for r in sorted(records, key=lambda x: (int(x['chassis'][1:]), int(x['slot']) if str(x['slot']).isdigit() else 99, int(x['point']) if str(x['point']).isdigit() else 999)):
        status = "SPARE" if r['is_spare'] else "ACTIVE"
        row_vals = [
            r['chassis'],
            r['slot_str'],
            r['point'],
            r['terminal'],
            r['card_model'],
            r['io_type'],
            r['tag'],
            r['desc'],
            r['dest'],
            r['jb_loc'] if r['jb_loc'] else JB_INFO.get(r['dest'], {}).get('loc', '-'),
            r['floor'],
            r['pid'],
            status
        ]
        ws_full.append(row_vals)
        ws_full.row_dimensions[f_row].height = 19
        fill_col = LIGHT_BG_1 if f_row % 2 == 0 else LIGHT_BG_2
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws_full.cell(row=f_row, column=col_idx)
            cell.font = Font(name="Arial", size=9, bold=(col_idx == 1))
            cell.fill = PatternFill(start_color=fill_col, end_color=fill_col, fill_type="solid")
            cell.border = thin_border()
            if col_idx in [1, 2, 3, 4, 6, 9, 11, 13]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
                
            # Status styling
            if col_idx == 13:
                if status == "ACTIVE":
                    cell.fill = PatternFill(start_color=ACTIVE_FILL, end_color=ACTIVE_FILL, fill_type="solid")
                    cell.font = Font(name="Arial", size=9, bold=True, color=ACTIVE_FONT)
                else:
                    cell.fill = PatternFill(start_color=SPARE_FILL, end_color=SPARE_FILL, fill_type="solid")
                    cell.font = Font(name="Arial", size=9, bold=False, color=SPARE_FONT)
        f_row += 1
        
    # Auto-fit column widths across all sheets
    for ws in [ws_dash, ws_sched, ws_mat, ws_full]:
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                # skip merged title rows
                if cell.row in [1, 2]: continue
                v = str(cell.value or '')
                if len(v) > max_len:
                    max_len = len(v)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)
            
    ws_dash.column_dimensions['A'].width = 14
    ws_dash.column_dimensions['B'].width = 34
    ws_dash.column_dimensions['C'].width = 38
    ws_sched.column_dimensions['K'].width = 45
    ws_full.column_dimensions['H'].width = 35
    
    wb.save(OUTPUT_EXCEL)
    print(f"Report successfully saved to: {OUTPUT_EXCEL}")

if __name__ == "__main__":
    build_report()
