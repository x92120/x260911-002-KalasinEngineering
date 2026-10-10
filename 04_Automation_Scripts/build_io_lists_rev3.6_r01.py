#!/usr/bin/env python3
"""
=============================================================================
KALASIN ENGINEERING xCIP-1545 AUTOMATION SYSTEM (Ref: x2608003)
PROJECT SPRINT 18K TPA SPRAY DRYER / JET COOKER - INGREDION (THAILAND) CO., LTD.
MASTER REVISION BUILDER: I/O LIST BY JUNCTION BOX AND SLOT CONFIG (-rev3.6-r01)
=============================================================================
Input Files (Folder: 03-IO_List):
  1. Instrument I-O List Rev.3.6a.xlsx (Client Master Instrument Schedule Rev 3.6a)
  2. IO_List-By_SlotConfig_rev02.xlsx (Standardized 13-Slot Architecture C1S0 - C5S4)

Output Files (Folder: 03-IO_List & 03_IO_Lists_and_Schedules/IO_List):
  - IO_List-By_SlotConfig-rev3.6-r01.xlsx
  - IO_List_By_SlotConfig-rev3.6-r01.xlsx (Alias)
  - IO_List_By_Junction_Box-rev3.6-r01.xlsx
  - Junction_Box_IO_List-rev3.6-r01.xlsx (Alias)
  - Master_IO_Mapping_Report-rev3.6-r01.xlsx
  - IO_Mapping_Report-rev3.6-r01.xlsx (Alias)
=============================================================================
"""

import os
import shutil
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.abspath(os.path.join(SCRIPT_DIR, ".."))
IO_DIR = os.path.join(BASE_DIR, "03-IO_List")
IO_DIR_ALT = os.path.join(BASE_DIR, "03_IO_Lists_and_Schedules", "IO_List")

INST_FILE = os.path.join(IO_DIR, "Instrument I-O List Rev.3.6a.xlsx")
SLOT_FILE = os.path.join(IO_DIR, "IO_List-By_SlotConfig_rev02.xlsx")

# Destination Output File Names (-rev3.6-r01)
OUT_SLOT_CONFIG = os.path.join(IO_DIR, "IO_List-By_SlotConfig-rev3.6-r01.xlsx")
OUT_SLOT_CONFIG_ALIAS = os.path.join(IO_DIR, "IO_List_By_SlotConfig-rev3.6-r01.xlsx")

OUT_JB = os.path.join(IO_DIR, "IO_List_By_Junction_Box-rev3.6-r01.xlsx")
OUT_JB_ALIAS = os.path.join(IO_DIR, "Junction_Box_IO_List-rev3.6-r01.xlsx")

OUT_MAPPING = os.path.join(IO_DIR, "Master_IO_Mapping_Report-rev3.6-r01.xlsx")
OUT_MAPPING_ALIAS = os.path.join(IO_DIR, "IO_Mapping_Report-rev3.6-r01.xlsx")

def load_inst_dict():
    print(f"Loading Client Master Instrument List: {INST_FILE}")
    wb_inst = openpyxl.load_workbook(INST_FILE, data_only=True)
    ws_inst = wb_inst['Rev.3']

    inst_lookup = {}
    for r in range(12, ws_inst.max_row + 1):
        item_no = ws_inst.cell(r, 1).value
        tag = ws_inst.cell(r, 2).value
        new_tag = ws_inst.cell(r, 3).value
        pid = ws_inst.cell(r, 4).value
        desc = ws_inst.cell(r, 5).value
        inst_name = ws_inst.cell(r, 6).value
        do_sig = ws_inst.cell(r, 7).value
        di_sig = ws_inst.cell(r, 8).value
        ao_sig = ws_inst.cell(r, 9).value
        ai_sig = ws_inst.cell(r, 10).value
        bus_sig = ws_inst.cell(r, 11).value
        sig_type = ws_inst.cell(r, 12).value
        sig_to = ws_inst.cell(r, 13).value
        rng = ws_inst.cell(r, 14).value
        fn = ws_inst.cell(r, 15).value
        cable = ws_inst.cell(r, 16).value
        prot = ws_inst.cell(r, 17).value
        remark = ws_inst.cell(r, 18).value

        if tag or new_tag or desc:
            item = {
                'item_no': str(item_no).strip() if item_no is not None else '',
                'tag': str(tag).strip() if tag else '',
                'new_tag': str(new_tag).strip() if new_tag else '',
                'pid': str(pid).strip() if pid else '',
                'desc': str(desc).strip() if desc else '',
                'inst_name': str(inst_name).strip() if inst_name else '',
                'sig_type': str(sig_type).strip() if sig_type else '',
                'sig_to': str(sig_to).strip() if sig_to else '',
                'range': str(rng).strip() if rng else '',
                'cable': str(cable).strip() if cable else '',
                'prot': str(prot).strip() if prot else '',
                'remark': str(remark).strip() if remark else ''
            }
            active_t = item['new_tag'] if item['new_tag'] else item['tag']
            item['active_tag'] = active_t

            for k in [active_t, item['new_tag'], item['tag']]:
                if k:
                    u = k.upper()
                    inst_lookup[u] = item
                    inst_lookup[u.replace(' ', '').replace('-', '')] = item

    print(f"Loaded {len(inst_lookup)} instrument keys from Rev.3.6a.")
    return inst_lookup

def build_updated_slot_config(inst_lookup):
    print(f"\nBuilding Updated Slot Config: {OUT_SLOT_CONFIG}")
    wb_slot = openpyxl.load_workbook(SLOT_FILE)

    # Update 00_Slot_Index header banner
    ws_idx = wb_slot['00_Slot_Index']
    ws_idx['A1'] = "KALASIN ENGINEERING xCIP-1545 AUTOMATION SYSTEM (Ref: x2608003 - Rev.3.6a-R01)"
    ws_idx['A2'] = "CHASSIS C1 - C4 (13-SLOT STANDARD) & C5: MASTER SLOT CONFIGURATION DIRECTORY (Rev.3.6a-R01)"

    slot_sheets = [s for s in wb_slot.sheetnames if s not in ['00_Slot_Index', 'C1SlotConfig', 'C2SlotConfig', 'C3SlotConfig', 'C4SlotConfig', 'C5SlotConfig']]

    updated_rows = 0
    for sname in slot_sheets:
        ws = wb_slot[sname]
        # Update title banner if present
        if ws['A1'].value and 'SLOT CONFIGURATION' in str(ws['A1'].value):
            if '- Rev.3.6a-R01' not in str(ws['A1'].value):
                ws['A1'].value = str(ws['A1'].value) + " - Rev.3.6a-R01"

        for r in range(6, ws.max_row + 1):
            t_no = ws.cell(r, 1).value
            if t_no is None: continue

            inst_tag = ws.cell(r, 8).value
            inst_desc = ws.cell(r, 9).value

            t_str = str(inst_tag or '').strip()
            if t_str and t_str not in ['Spare', 'None', '-']:
                u_tag = t_str.upper()
                u_clean = u_tag.replace(' ', '').replace('-', '')
                inst = inst_lookup.get(u_tag) or inst_lookup.get(u_clean)
                if inst:
                    new_t = inst['active_tag']
                    new_d = inst['desc']
                    if new_t and new_t != t_str:
                        ws.cell(r, 8).value = new_t
                    if new_d and new_d != str(inst_desc or '').strip():
                        ws.cell(r, 9).value = new_d
                    updated_rows += 1

    print(f"Updated {updated_rows} channel rows with Rev.3.6a instrument data.")
    wb_slot.save(OUT_SLOT_CONFIG)
    shutil.copyfile(OUT_SLOT_CONFIG, OUT_SLOT_CONFIG_ALIAS)
    print(f"Saved: {OUT_SLOT_CONFIG}")
    print(f"Saved: {OUT_SLOT_CONFIG_ALIAS}")

    if os.path.exists(IO_DIR_ALT):
        alt_dest = os.path.join(IO_DIR_ALT, os.path.basename(OUT_SLOT_CONFIG))
        shutil.copyfile(OUT_SLOT_CONFIG, alt_dest)
        print(f"Saved: {alt_dest}")

def build_junction_box_report():
    print(f"\nBuilding Junction Box I/O List Report (-rev3.6-r01)...")
    import build_io_list_by_junction_box as jb_builder

    # Monkeypatch output paths
    jb_builder.OUTPUT_FILE = OUT_JB
    jb_builder.OUTPUT_FILE_ALT = OUT_JB_ALIAS
    jb_builder.SLOT_FILE = OUT_SLOT_CONFIG

    jb_builder.main()

    if os.path.exists(IO_DIR_ALT):
        alt_dest = os.path.join(IO_DIR_ALT, os.path.basename(OUT_JB))
        shutil.copyfile(OUT_JB, alt_dest)
        print(f"Saved: {alt_dest}")

def build_master_mapping_report():
    print(f"\nBuilding Master I/O Mapping Report (-rev3.6-r01)...")
    import build_master_io_mapping_report as map_builder

    # Monkeypatch output paths
    map_builder.OUTPUT_FILE = OUT_MAPPING_ALIAS
    map_builder.OUTPUT_FILE_ALT = OUT_MAPPING
    map_builder.SLOT_FILE = OUT_SLOT_CONFIG

    map_builder.main()

    if os.path.exists(IO_DIR_ALT):
        alt_dest = os.path.join(IO_DIR_ALT, os.path.basename(OUT_MAPPING))
        shutil.copyfile(OUT_MAPPING, alt_dest)
        print(f"Saved: {alt_dest}")

def main():
    print("=================================================================")
    print("MASTER REVISION BUILDER: I/O LIST BY JUNCTION BOX & SLOT CONFIG")
    print("REVISION TAG: -rev3.6-r01")
    print("=================================================================")

    inst_lookup = load_inst_dict()
    build_updated_slot_config(inst_lookup)
    build_junction_box_report()
    build_master_mapping_report()

    print("\n=================================================================")
    print("ALL REVISION -rev3.6-r01 I/O LIST WORKBOOKS SUCCESSFULLY BUILT!")
    print("=================================================================")

if __name__ == '__main__':
    main()
