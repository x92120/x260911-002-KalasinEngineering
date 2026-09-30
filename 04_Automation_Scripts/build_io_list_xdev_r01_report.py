import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SOURCE_EXCEL = os.path.join(BASE_DIR, "03_IO_Lists_and_Schedules", "IO_List_xDev-R01-Tag35-7-2.xlsx")
REPORT_EXCEL = os.path.join(BASE_DIR, "03_IO_Lists_and_Schedules", "IO_List_xDev-R01_Report.xlsx")

def build_io_report():
    print(f"Loading source workbook: {SOURCE_EXCEL}")
    wb_src = openpyxl.load_workbook(SOURCE_EXCEL, data_only=True)
    sheet_src = wb_src['IO List']

    headers_src = {sheet_src.cell(1, c).value.strip(): c for c in range(1, sheet_src.max_column + 1) if sheet_src.cell(1, c).value}

    chassis_col = headers_src['Chassis']
    dest_col = headers_src['Destination']
    io_type_col = headers_src['I/O Type']
    jb_loc_col = headers_src.get('Junction Box Location')

    # Aggregations
    matrix_counts = defaultdict(int) # (chassis, io_type, destination_canonical)
    jb_summary = defaultdict(lambda: {'location': '', 'DI': 0, 'DO': 0, 'AI': 0, 'AO': 0, 'BUS': 0, 'TOTAL': 0})
    chassis_summary = defaultdict(lambda: {'DI': 0, 'DO': 0, 'AI': 0, 'AO': 0, 'BUS': 0, 'TOTAL': 0})

    dest_canonical_map = {
        'JB-401': 'JB401', 'JB401': 'JB401',
        'JB-402': 'JB402', 'JB402': 'JB402',
        'JB-601': 'JB601', 'JB601': 'JB601',
        'JB-602': 'JB602', 'JB602': 'JB602',
        'JB-606': 'JB606', 'JB606': 'JB606',
        'JB-607': 'JB607', 'JB607': 'JB607',
        'JB-608': 'JB608', 'JB608': 'JB608',
        'JB-612': 'JB612', 'JB612': 'JB612',
        'JB-618': 'JB618', 'JB618': 'JB618',
        'IS-JB-603': 'IS-JB-603',
        'IS-JB-608': 'IS-JB-608',
        'IS-JB-612': 'IS-JB-612',
        'IS-JB-618': 'IS-JB-618',
        'RIO-200': 'RIO-200',
        'CA1': 'CA1',
        'MCC': 'MCC'
    }

    total_signals = 0

    for r in range(2, sheet_src.max_row + 1):
        ch = str(sheet_src.cell(r, chassis_col).value or '').strip()
        dest = str(sheet_src.cell(r, dest_col).value or '').strip()
        io_t = str(sheet_src.cell(r, io_type_col).value or '').strip()
        jb_loc = str(sheet_src.cell(r, jb_loc_col).value or '').strip() if jb_loc_col else ''

        if ch or dest or io_t:
            total_signals += 1
            canon_dest = dest_canonical_map.get(dest, dest)

            matrix_counts[(ch, io_t, canon_dest)] += 1

            if io_t in ['DI', 'DO', 'AI', 'AO', 'BUS']:
                jb_summary[canon_dest][io_t] += 1
                chassis_summary[ch][io_t] += 1
            else:
                # Group ETH/CPU/135 appropriately if needed
                pass
            
            jb_summary[canon_dest]['TOTAL'] += 1
            if jb_loc and not jb_summary[canon_dest]['location']:
                jb_summary[canon_dest]['location'] = jb_loc
            
            chassis_summary[ch]['TOTAL'] += 1

    print(f"Processed {total_signals} total I/O records.")

    # Create Report Workbook
    wb_rep = openpyxl.Workbook()
    ws = wb_rep.active
    ws.title = "IO_Summary_Matrix"
    ws.views.sheetView[0].showGridLines = True

    # Styling Palette
    FONT_FAMILY = "Segoe UI"
    NAVY_FILL = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    STEEL_FILL = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    ACCENT_FILL = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    ZEBRA_FILL = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    TOTAL_FILL = PatternFill(start_color="F2F4F7", end_color="F2F4F7", fill_type="solid")
    NONZERO_FILL = PatternFill(start_color="EBF1F5", end_color="EBF1F5", fill_type="solid")

    TITLE_FONT = Font(name=FONT_FAMILY, size=16, bold=True, color="FFFFFF")
    SUBTITLE_FONT = Font(name=FONT_FAMILY, size=10, italic=True, color="DCE6F1")
    HEADER_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")
    SUBHEADER_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="1F497D")
    BOLD_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="000000")
    REGULAR_FONT = Font(name=FONT_FAMILY, size=10, color="000000")
    MUTED_FONT = Font(name=FONT_FAMILY, size=10, color="A6A6A6")

    THIN_BORDER_SIDE = Side(border_style="thin", color="D9D9D9")
    THICK_BOTTOM_SIDE = Side(border_style="medium", color="1F497D")
    DOUBLE_BOTTOM_SIDE = Side(border_style="double", color="1F497D")

    CELL_BORDER = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=THIN_BORDER_SIDE)
    HEADER_BORDER = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=THICK_BOTTOM_SIDE)
    TOTAL_BORDER = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=DOUBLE_BOTTOM_SIDE)

    # 1. Title Block
    ws.merge_cells("A1:Q1")
    ws["A1"] = "KALASIN ENGINEERING - PLC I/O SYSTEM SUMMARY REPORT"
    ws["A1"].font = TITLE_FONT
    ws["A1"].fill = NAVY_FILL
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:Q2")
    ws["A2"] = "Chassis to Junction Box & Enclosure Signal Distribution Matrix (Source: IO_List_xDev-R01-Tag35-7-2.xlsx)"
    ws["A2"].font = SUBTITLE_FONT
    ws["A2"].fill = NAVY_FILL
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 20

    # 2. Main Matrix Headers (Row 3)
    dest_headers = [
        'JB401', 'JB402', 'JB601', 'JB602', 'JB606', 'JB607', 'JB608', 'JB612', 'JB618',
        'IS-JB-603', 'IS-JB-608', 'IS-JB-612', 'IS-JB-618',
        'RIO-200', 'CA1', 'MCC'
    ]

    ws.cell(3, 1, "Chassis").font = HEADER_FONT
    ws.cell(3, 1).fill = STEEL_FILL
    ws.cell(3, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(3, 1).border = HEADER_BORDER

    ws.cell(3, 2, "Signal").font = HEADER_FONT
    ws.cell(3, 2).fill = STEEL_FILL
    ws.cell(3, 2).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(3, 2).border = HEADER_BORDER

    for col_idx, jb_name in enumerate(dest_headers, start=3):
        cell = ws.cell(3, col_idx, jb_name)
        cell.font = HEADER_FONT
        cell.fill = STEEL_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = HEADER_BORDER

    total_col_idx = len(dest_headers) + 3
    tot_cell = ws.cell(3, total_col_idx, "TOTAL")
    tot_cell.font = HEADER_FONT
    tot_cell.fill = NAVY_FILL
    tot_cell.alignment = Alignment(horizontal="center", vertical="center")
    tot_cell.border = HEADER_BORDER

    ws.row_dimensions[3].height = 24

    # Rows definition
    chassis_list = ['C1', 'C2', 'C3', 'C4', 'C5', 'C8', 'C99']
    signal_types = ['DI', 'DO', 'AI', 'AO']

    curr_row = 4

    for ch in chassis_list:
        ch_start_row = curr_row
        for sig in signal_types:
            ws.cell(curr_row, 1, ch).font = BOLD_FONT
            ws.cell(curr_row, 1).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(curr_row, 1).border = CELL_BORDER

            ws.cell(curr_row, 2, sig).font = BOLD_FONT
            ws.cell(curr_row, 2).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(curr_row, 2).border = CELL_BORDER

            row_total = 0
            for c_idx, jb_name in enumerate(dest_headers, start=3):
                cnt = matrix_counts.get((ch, sig, jb_name), 0)
                cell = ws.cell(curr_row, c_idx)
                if cnt > 0:
                    cell.value = cnt
                    cell.font = BOLD_FONT
                    cell.fill = NONZERO_FILL
                else:
                    cell.value = ""
                    cell.font = MUTED_FONT

                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = CELL_BORDER
                row_total += cnt

            tot_cell = ws.cell(curr_row, total_col_idx)
            tot_cell.value = row_total if row_total > 0 else ""
            tot_cell.font = BOLD_FONT
            tot_cell.fill = ACCENT_FILL
            tot_cell.alignment = Alignment(horizontal="center", vertical="center")
            tot_cell.border = CELL_BORDER

            ws.row_dimensions[curr_row].height = 19
            curr_row += 1

        # Merge Chassis column for the 4 rows
        ws.merge_cells(start_row=ch_start_row, start_column=1, end_row=curr_row - 1, end_column=1)

    # 3. Summary Total Row
    ws.cell(curr_row, 1, "SYSTEM").font = TITLE_FONT
    ws.cell(curr_row, 1).fill = NAVY_FILL
    ws.cell(curr_row, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(curr_row, 1).border = TOTAL_BORDER

    ws.cell(curr_row, 2, "TOTAL").font = HEADER_FONT
    ws.cell(curr_row, 2).fill = NAVY_FILL
    ws.cell(curr_row, 2).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(curr_row, 2).border = TOTAL_BORDER

    grand_total = 0
    for c_idx, jb_name in enumerate(dest_headers, start=3):
        col_letter = get_column_letter(c_idx)
        sum_formula = f"=SUM({col_letter}4:{col_letter}{curr_row-1})"
        cell = ws.cell(curr_row, c_idx, sum_formula)
        cell.font = HEADER_FONT
        cell.fill = NAVY_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = TOTAL_BORDER

    tot_col_letter = get_column_letter(total_col_idx)
    grand_formula = f"=SUM({tot_col_letter}4:{tot_col_letter}{curr_row-1})"
    cell = ws.cell(curr_row, total_col_idx, grand_formula)
    cell.font = TITLE_FONT
    cell.fill = NAVY_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = TOTAL_BORDER

    ws.row_dimensions[curr_row].height = 25

    # 4. Create Sheet 2: JB_Summary_Details
    ws2 = wb_rep.create_sheet(title="JB_Summary_Details")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:H1")
    ws2["A1"] = "JUNCTION BOX & ENCLOSURE CAPACITY SUMMARY"
    ws2["A1"].font = TITLE_FONT
    ws2["A1"].fill = NAVY_FILL
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 28

    headers2 = ["Enclosure Name", "Location / Description", "DI Points", "DO Points", "AI Points", "AO Points", "BUS Points", "Total I/O Points"]
    for c_idx, h_text in enumerate(headers2, start=1):
        cell = ws2.cell(2, c_idx, h_text)
        cell.font = HEADER_FONT
        cell.fill = STEEL_FILL
        cell.alignment = Alignment(horizontal="center" if c_idx != 2 else "left", vertical="center")
        cell.border = HEADER_BORDER

    ws2.row_dimensions[2].height = 24

    r2 = 3
    for jb_name in sorted(jb_summary.keys()):
        if not jb_name or jb_name == '135':
            continue
        info = jb_summary[jb_name]
        ws2.cell(r2, 1, jb_name).font = BOLD_FONT
        ws2.cell(r2, 1).alignment = Alignment(horizontal="center", vertical="center")
        ws2.cell(r2, 1).border = CELL_BORDER

        ws2.cell(r2, 2, info['location'] or "Field Enclosure").font = REGULAR_FONT
        ws2.cell(r2, 2).alignment = Alignment(horizontal="left", vertical="center")
        ws2.cell(r2, 2).border = CELL_BORDER

        ws2.cell(r2, 3, info['DI'] or 0).font = REGULAR_FONT
        ws2.cell(r2, 3).alignment = Alignment(horizontal="right", vertical="center")
        ws2.cell(r2, 3).border = CELL_BORDER

        ws2.cell(r2, 4, info['DO'] or 0).font = REGULAR_FONT
        ws2.cell(r2, 4).alignment = Alignment(horizontal="right", vertical="center")
        ws2.cell(r2, 4).border = CELL_BORDER

        ws2.cell(r2, 5, info['AI'] or 0).font = REGULAR_FONT
        ws2.cell(r2, 5).alignment = Alignment(horizontal="right", vertical="center")
        ws2.cell(r2, 5).border = CELL_BORDER

        ws2.cell(r2, 6, info['AO'] or 0).font = REGULAR_FONT
        ws2.cell(r2, 6).alignment = Alignment(horizontal="right", vertical="center")
        ws2.cell(r2, 6).border = CELL_BORDER

        ws2.cell(r2, 7, info['BUS'] or 0).font = REGULAR_FONT
        ws2.cell(r2, 7).alignment = Alignment(horizontal="right", vertical="center")
        ws2.cell(r2, 7).border = CELL_BORDER

        ws2.cell(r2, 8, f"=SUM(C{r2}:G{r2})").font = BOLD_FONT
        ws2.cell(r2, 8).alignment = Alignment(horizontal="right", vertical="center")
        ws2.cell(r2, 8).border = CELL_BORDER
        ws2.cell(r2, 8).fill = ACCENT_FILL

        ws2.row_dimensions[r2].height = 20
        r2 += 1

    # Total Row for Sheet 2
    ws2.cell(r2, 1, "TOTAL").font = HEADER_FONT
    ws2.cell(r2, 1).fill = NAVY_FILL
    ws2.cell(r2, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws2.cell(r2, 1).border = TOTAL_BORDER

    ws2.cell(r2, 2, "").font = HEADER_FONT
    ws2.cell(r2, 2).fill = NAVY_FILL
    ws2.cell(r2, 2).border = TOTAL_BORDER

    for c_idx in range(3, 9):
        col_letter = get_column_letter(c_idx)
        cell = ws2.cell(r2, c_idx, f"=SUM({col_letter}3:{col_letter}{r2-1})")
        cell.font = HEADER_FONT
        cell.fill = NAVY_FILL
        cell.alignment = Alignment(horizontal="right", vertical="center")
        cell.border = TOTAL_BORDER

    ws2.row_dimensions[r2].height = 24

    # Adjust Column Widths
    for sheet in [ws, ws2]:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    print(f"Saving updated report to: {REPORT_EXCEL}")
    wb_rep.save(REPORT_EXCEL)
    print("Report generated successfully!")

if __name__ == "__main__":
    build_io_report()
