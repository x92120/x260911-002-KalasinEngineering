#!/usr/bin/env python3
"""
================================================================================
Kalasin Engineering xCIP System (Project Ref: x2608003)
CHASSIS C1 DESTINATION ROUTING & I/O SUMMARY PDF GENERATOR (TYPST)
================================================================================
Generates:
  1. 03_IO_Lists_and_Schedules/Chassis_C1_Destination_and_IO_Summary_Report.pdf
  2. 02_Electrical_and_eDrawing/04_PDF_Reports_and_Drawings/Chassis_C1_Destination_and_IO_Summary_Report.pdf
================================================================================
"""

import os
import subprocess
from collections import defaultdict, Counter
import openpyxl

BASE_DIR = "/Users/x92120/xApp-001/x260911-002-KalasinEngineering"
SOURCE_EXCEL = os.path.join(BASE_DIR, "03_IO_Lists_and_Schedules/IO_List_xDev-R01-Tag35-7.xlsx")
TYP_FILE = os.path.join(BASE_DIR, "04_Automation_Scripts/chassis_c1_report.typ")
PDF_OUT_1 = os.path.join(BASE_DIR, "03_IO_Lists_and_Schedules/Chassis_C1_Destination_and_IO_Summary_Report.pdf")
PDF_OUT_2 = os.path.join(BASE_DIR, "02_Electrical_and_eDrawing/04_PDF_Reports_and_Drawings/Chassis_C1_Destination_and_IO_Summary_Report.pdf")
TYPST_BIN = "/opt/homebrew/bin/typst"

os.makedirs(os.path.dirname(PDF_OUT_2), exist_ok=True)

def load_c1_data():
    wb = openpyxl.load_workbook(SOURCE_EXCEL, data_only=True)
    ws = wb['IO List']
    
    c1_records = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        col7 = str(r[2]).strip() if r[2] is not None else ''
        ch = str(r[30]).strip() if r[30] is not None else ''
        slot = str(r[31]).strip() if r[31] is not None else ''
        if not ch and col7.startswith('C') and 'S' in col7:
            ch = col7.split('S')[0]
            slot = col7.split('S')[1].split(',')[0]
            
        if ch == 'C1':
            dest = str(r[1]).strip() if r[1] is not None else ''
            card = str(r[3]).strip() if r[3] is not None else ''
            card_type = str(r[26]).strip() if r[26] is not None else ''
            term = str(r[5]).strip() if r[5] is not None else ''
            tag = str(r[11]).strip() if r[11] is not None else ''
            desc = str(r[12]).strip() if r[12] is not None else ''
            io_t = str(r[20]).strip() if r[20] is not None else ''
            pt = str(r[32]).strip() if r[32] is not None else ''
            pid = str(r[14]).strip() if r[14] is not None else ''
            
            # Clean card
            c_name = card if card and card != '#N/A' else card_type
            if '-' in c_name:
                parts = c_name.split('-')
                if len(parts) >= 2 and not parts[1].isdigit():
                    clean_c = f"{parts[0]}-{parts[1]}"
                else:
                    clean_c = c_name
            else:
                clean_c = c_name
                
            is_sp = 'SPARE' in tag.upper() or not tag or tag == '#N/A' or tag == '0'
            c1_records.append({
                'slot': int(slot) if slot.isdigit() else slot,
                'pt': int(pt) if str(pt).isdigit() else pt,
                'dest': dest,
                'card': clean_c,
                'term': term,
                'tag': tag if not is_sp else 'SPARE',
                'desc': desc if (desc and desc != '#N/A' and desc != '0') else ('Spare Reserve Channel' if is_sp else '-'),
                'io_t': io_t,
                'pid': pid if pid != '#N/A' else '-',
                'is_sp': is_sp
            })
            
    c1_records.sort(key=lambda x: (x['slot'], x['pt']))
    return c1_records

def escape_typst(s):
    if not s: return ""
    return str(s).replace('\\', '\\\\').replace('"', '\\"').replace('$', '\\$').replace('#', '\\#').replace('_', '\\_')

def generate_typst_content(records):
    # Destination summary stats
    dest_stats = defaultdict(lambda: {'total': 0, 'active': 0, 'spare': 0, 'DI': 0, 'DO': 0, 'AI': 0, 'Comm': 0})
    for r in records:
        d = r['dest']
        dest_stats[d]['total'] += 1
        if r['is_sp']: dest_stats[d]['spare'] += 1
        else: dest_stats[d]['active'] += 1
        
        io = r['io_t']
        if io in ['DI', 'DO', 'AI']:
            dest_stats[d][io] += 1
        else:
            dest_stats[d]['Comm'] += 1

    # Slot summary stats
    slot_stats = defaultdict(lambda: defaultdict(lambda: {'total': 0, 'active': 0, 'spare': 0, 'card': '', 'io_t': '', 'tags': []}))
    for r in records:
        s = r['slot']
        d = r['dest']
        st = slot_stats[s][d]
        st['total'] += 1
        st['card'] = r['card']
        st['io_t'] = r['io_t']
        if r['is_sp']: st['spare'] += 1
        else:
            st['active'] += 1
            if r['tag'] not in st['tags']: st['tags'].append(r['tag'])

    typ = []
    typ.append("""
#set page(
  paper: "a4",
  flipped: true,
  margin: (top: 1.2cm, bottom: 1.2cm, left: 1.4cm, right: 1.4cm),
  header: context {
    grid(
      columns: (1fr, 1fr),
      align: (left, right),
      text(size: 8pt, fill: rgb("64748b"), font: "Arial")[
        *KALASIN ENGINEERING* | Turnkey xCIP-1545 Automation System (Ref: x2608003)
      ],
      text(size: 8pt, fill: rgb("64748b"), font: "Arial")[
        *Chassis C1 Destination & I/O Summary Report* | Page #counter(page).display()
      ]
    )
    line(length: 100%, stroke: 0.5pt + rgb("cbd5e1"))
  },
  footer: context {
    line(length: 100%, stroke: 0.5pt + rgb("cbd5e1"))
    v(2pt)
    grid(
      columns: (1fr, 1fr),
      align: (left, right),
      text(size: 8pt, fill: rgb("94a3b8"), font: "Arial")[
        Client: Ingredion (Thailand) Co., Ltd. / ISHITON | Document Ref: KAL-IO-REP-C1-001
      ],
      text(size: 8pt, fill: rgb("94a3b8"), font: "Arial")[
        Source File: `IO_List_xDev-R01-Tag35-7.xlsx` | Approved Rev: 3.6
      ]
    )
  }
)

#set text(font: "Arial", size: 8.5pt, fill: rgb("0f172a"))
#set par(justify: true, leading: 0.5em)

// --- TITLE BLOCK ---
#block(
  fill: rgb("1e293b"),
  inset: (x: 16pt, y: 11pt),
  radius: 4pt,
  width: 100%
)[
  #grid(
    columns: (2.5fr, 1fr),
    align: (left + horizon, right + horizon),
    [
      #text(fill: white, size: 14pt, weight: "bold")[
        CHASSIS C1: DESTINATION ROUTING & I/O POINT SUMMARY
      ]
      #v(2pt)
      #text(fill: rgb("93c5fd"), size: 9pt, weight: "medium")[
        ROCKWELL 1756 CONTROLLOGIX PROCESSOR RACK (CA1) --- POINT COUNT & FIELD ROUTING SCHEDULE
      ]
    ],
    [
      #text(fill: white, size: 8pt)[
        *Project:* Kalasin xCIP Turnkey #linebreak()
        *Controller:* 1756-L950TPSXT #linebreak()
        *Rack Enclosure:* Control Cabinet CA1
      ]
    ]
  )
]

#v(8pt)

// --- SECTION 1: EXECUTIVE KPI SUMMARY ---
#text(size: 11pt, weight: "bold", fill: rgb("1e3a8a"))[1. Executive Summary & Routing Breakdown]
#v(4pt)

#grid(
  columns: (1fr, 1fr, 1fr, 1fr),
  gutter: 10pt,
  block(fill: rgb("f1f5f9"), stroke: 1pt + rgb("cbd5e1"), inset: 8pt, radius: 3pt)[
    #text(size: 8pt, fill: rgb("64748b"))[*TOTAL C1 CHANNELS*] #linebreak()
    #text(size: 16pt, weight: "bold", fill: rgb("0f172a"))[228 Points] #linebreak()
    #text(size: 7.5pt, fill: rgb("475569"))[Configured across Slots 00--11]
  ],
  block(fill: rgb("dcfce7"), stroke: 1pt + rgb("86efac"), inset: 8pt, radius: 3pt)[
    #text(size: 8pt, fill: rgb("166534"))[*ACTIVE SIGNALS*] #linebreak()
    #text(size: 16pt, weight: "bold", fill: rgb("14532d"))[138 Points] #linebreak()
    #text(size: 7.5pt, fill: rgb("15803d"))[60.5% Active Field Instruments]
  ],
  block(fill: rgb("fef9c3"), stroke: 1pt + rgb("fde047"), inset: 8pt, radius: 3pt)[
    #text(size: 8pt, fill: rgb("854d0e"))[*SPARE CAPACITY*] #linebreak()
    #text(size: 16pt, weight: "bold", fill: rgb("713f12"))[90 Points] #linebreak()
    #text(size: 7.5pt, fill: rgb("a16207"))[39.5% Expansion Reserve]
  ],
  block(fill: rgb("e0e7ff"), stroke: 1pt + rgb("a5b4fc"), inset: 8pt, radius: 3pt)[
    #text(size: 8pt, fill: rgb("3730a3"))[*DESTINATIONS*] #linebreak()
    #text(size: 16pt, weight: "bold", fill: rgb("1e1b4b"))[3 Locations] #linebreak()
    #text(size: 7.5pt, fill: rgb("4338ca"))[CA1, JB-401, JB-601]
  ]
)

#v(8pt)

// --- TABLE 1: DESTINATION SUMMARY ---
#table(
  columns: (1.5fr, 2.5fr, 1fr, 1fr, 1fr, 1fr, 1fr, 1fr, 1fr),
  fill: (col, row) => if row == 0 { rgb("1e293b") } else if calc.even(row) { rgb("f8fafc") } else { rgb("ffffff") },
  stroke: (x, y) => if y == 0 { 0.5pt + rgb("475569") } else { 0.5pt + rgb("e2e8f0") },
  align: (col, row) => if col < 2 { left + horizon } else { center + horizon },
  inset: 5pt,
  
  // Headers
  text(fill: white, weight: "bold", size: 8pt)[Destination],
  text(fill: white, weight: "bold", size: 8pt)[Plant Location / Role],
  text(fill: white, weight: "bold", size: 8pt)[Total Pts],
  text(fill: white, weight: "bold", size: 8pt)[Active],
  text(fill: white, weight: "bold", size: 8pt)[Spare],
  text(fill: white, weight: "bold", size: 8pt)[DI (24V)],
  text(fill: white, weight: "bold", size: 8pt)[DO (24V)],
  text(fill: white, weight: "bold", size: 8pt)[AI (4-20mA)],
  text(fill: white, weight: "bold", size: 8pt)[CPU / Comm],

  // Row 1: CA1
  [*CA1*], [Main Control Room (Local Cabinet Spares & Network)], [36], [0], [36], [16], [16], [0], [4],
  // Row 2: JB-401
  [*JB-401*], [Infeed 2nd Floor (Raw Material Prep & Cooking Infeed)], [112], [85], [27], [64], [32], [16], [0],
  // Row 3: JB-601
  [*JB-601*], [Spray Dryer 1st Floor (Chamber Bottom Discharge & Pumps)], [80], [53], [27], [48], [16], [16], [0],
  
  // Total Row
  table.cell(fill: rgb("334155"))[#text(fill: white, weight: "bold")[TOTAL]],
  table.cell(fill: rgb("334155"))[#text(fill: white, weight: "bold")[Chassis C1 Master Total]],
  table.cell(fill: rgb("334155"))[#text(fill: white, weight: "bold")[228]],
  table.cell(fill: rgb("334155"))[#text(fill: white, weight: "bold")[138]],
  table.cell(fill: rgb("334155"))[#text(fill: white, weight: "bold")[90]],
  table.cell(fill: rgb("334155"))[#text(fill: white, weight: "bold")[128]],
  table.cell(fill: rgb("334155"))[#text(fill: white, weight: "bold")[64]],
  table.cell(fill: rgb("334155"))[#text(fill: white, weight: "bold")[32]],
  table.cell(fill: rgb("334155"))[#text(fill: white, weight: "bold")[4]]
)

#v(10pt)

// --- SECTION 2: SLOT-BY-SLOT SCHEDULE ---
#text(size: 11pt, weight: "bold", fill: rgb("1e3a8a"))[2. Chassis C1 Slot-to-Destination Routing Schedule]
#v(4pt)

#table(
  columns: (0.7fr, 1.4fr, 0.9fr, 1fr, 0.8fr, 0.8fr, 0.8fr, 3.5fr),
  fill: (col, row) => if row == 0 { rgb("1e293b") } else if calc.even(row) { rgb("f8fafc") } else { rgb("ffffff") },
  stroke: (x, y) => if y == 0 { 0.5pt + rgb("475569") } else { 0.5pt + rgb("e2e8f0") },
  align: (col, row) => if col in (0, 2, 3, 4, 5, 6) { center + horizon } else { left + horizon },
  inset: 4.5pt,

  text(fill: white, weight: "bold", size: 8pt)[Slot],
  text(fill: white, weight: "bold", size: 8pt)[Hardware Card],
  text(fill: white, weight: "bold", size: 8pt)[Type],
  text(fill: white, weight: "bold", size: 8pt)[Destination],
  text(fill: white, weight: "bold", size: 8pt)[Points],
  text(fill: white, weight: "bold", size: 8pt)[Active],
  text(fill: white, weight: "bold", size: 8pt)[Spare],
  text(fill: white, weight: "bold", size: 8pt)[Connected Equipment / Purpose],

  // Slot 00
  [00], [1756-L950TPSXT], [CPU], [CA1], [1], [0], [1], [ControlLogix Main Controller Head],
  // Slot 01
  [01], [1756-EN4TR], [ETH-1], [CA1], [1], [0], [1], [Plant Network Port 1 (Enterprise / SCADA)],
  // Slot 02
  [02], [1756-EN4TR], [ETH-2], [CA1], [1], [0], [1], [DLR Device Level Ring Supervisor Node],
  // Slot 03
  [03], [1756-EN4TR], [ETH-3], [CA1], [1], [0], [1], [Historian & Auxiliary Communications Link],
  // Slot 04
  [04], [1756-IB32], [DI], [CA1], [16], [0], [16], [Local Cabinet Spare DI Bank (Channel 00--15)],
  [04], [1756-IB32], [DI], [JB-401], [16], [16], [0], [Infeed Flow & Limit Switches (FS-40101, FS-40102, ZS-40101)],
  // Slot 05
  [05], [1756-IB32], [DI], [JB-401], [32], [32], [0], [Infeed Level Switches (LSH/LSL-40101--08) & Valve Feedback],
  // Slot 06
  [06], [1756-IB32], [DI], [JB-401], [16], [2], [14], [Infeed Interlocks & Permissives],
  [06], [1756-IB32], [DI], [JB-601], [16], [10], [6], [Spray Dryer Chamber Bottom Level & Safety Switches],
  // Slot 07
  [07], [1756-IB32], [DI], [JB-601], [32], [23], [9], [Spray Dryer 1st Fl Flow & Pressure (FS-60101, PSH-60101)],
  // Slot 08
  [08], [1756-OB32], [DO], [CA1], [16], [0], [16], [Local Cabinet Spare DO Bank (Channel 00--15)],
  [08], [1756-OB32], [DO], [JB-401], [16], [16], [0], [Infeed Solenoid Valves (XV-40101--XV-40116)],
  // Slot 09
  [09], [1756-OB32], [DO], [JB-401], [16], [8], [8], [Infeed Process Valves (XV-40117--XV-40124)],
  [09], [1756-OB32], [DO], [JB-601], [16], [12], [4], [Spray Dryer Discharge Solenoids (XV-60101--XV-60112)],
  // Slot 10
  [10], [1756-IF16], [AI], [JB-401], [16], [11], [5], [Infeed Transmitters (PT-40101, TT-40101--04, FT-40101)],
  // Slot 11
  [11], [1756-IF16], [AI], [JB-601], [16], [8], [8], [Spray Dryer 1st Fl Pressure & Temp (PT-60101, TT-60101--06)]
)

#pagebreak()

// --- SECTION 3: COMPLETE POINT-BY-POINT SCHEDULE (228 CHANNELS) ---
#text(size: 11pt, weight: "bold", fill: rgb("1e3a8a"))[3. Chassis C1 Complete Point-by-Point Wiring & Loop Schedule (228 Channels)]
#v(4pt)

#table(
  columns: (0.6fr, 0.5fr, 1.2fr, 0.7fr, 1.4fr, 3.2fr, 1fr, 0.9fr, 0.8fr),
  fill: (col, row) => if row == 0 { rgb("1e293b") } else if calc.even(row) { rgb("f8fafc") } else { rgb("ffffff") },
  stroke: (x, y) => if y == 0 { 0.5pt + rgb("475569") } else { 0.5pt + rgb("e2e8f0") },
  align: (col, row) => if col in (0, 1, 3, 6, 7, 8) { center + horizon } else { left + horizon },
  inset: 3.5pt,

  text(fill: white, weight: "bold", size: 7.5pt)[Slot],
  text(fill: white, weight: "bold", size: 7.5pt)[Pt],
  text(fill: white, weight: "bold", size: 7.5pt)[Terminal],
  text(fill: white, weight: "bold", size: 7.5pt)[Type],
  text(fill: white, weight: "bold", size: 7.5pt)[Tag],
  text(fill: white, weight: "bold", size: 7.5pt)[Description],
  text(fill: white, weight: "bold", size: 7.5pt)[Dest JB],
  text(fill: white, weight: "bold", size: 7.5pt)[P&ID No.],
  text(fill: white, weight: "bold", size: 7.5pt)[Status],
""")

    for r in records:
        s_val = f"{r['slot']:02d}" if str(r['slot']).isdigit() else str(r['slot'])
        pt_val = f"{r['pt']:02d}" if str(r['pt']).isdigit() else str(r['pt'])
        term_val = escape_typst(r['term'])
        io_val = escape_typst(r['io_t'])
        tag_val = escape_typst(r['tag'])
        desc_val = escape_typst(r['desc'])
        dest_val = escape_typst(r['dest'])
        pid_val = escape_typst(r['pid'])
        status_val = "SPARE" if r['is_sp'] else "ACTIVE"
        
        status_cell = f'table.cell(fill: rgb("{"dcfce7" if status_val == "ACTIVE" else "fef9c3"}"))[#text(fill: rgb("{"14532d" if status_val == "ACTIVE" else "854d0e"}"), size: 7pt, weight: "{"bold" if status_val == "ACTIVE" else "regular"}")[{status_val}]]'
        
        typ.append(f'  [{s_val}], [{pt_val}], [{term_val}], [{io_val}], [{tag_val}], [{desc_val}], [{dest_val}], [{pid_val}], {status_cell},')

    typ.append(")\n")
    return "\n".join(typ)

def main():
    print(f"Loading data from {SOURCE_EXCEL}...")
    records = load_c1_data()
    print(f"Loaded {len(records)} C1 records.")
    
    print(f"Generating Typst document at: {TYP_FILE}")
    content = generate_typst_content(records)
    with open(TYP_FILE, "w", encoding="utf-8") as f:
        f.write(content)
        
    print(f"Compiling Typst to PDF: {PDF_OUT_1}")
    cmd1 = [TYPST_BIN, "compile", TYP_FILE, PDF_OUT_1]
    res1 = subprocess.run(cmd1, capture_output=True, text=True)
    if res1.returncode != 0:
        print(f"Error compiling Typst PDF 1: {res1.stderr}")
        return
        
    print(f"Compiling Typst to PDF: {PDF_OUT_2}")
    cmd2 = [TYPST_BIN, "compile", TYP_FILE, PDF_OUT_2]
    res2 = subprocess.run(cmd2, capture_output=True, text=True)
    if res2.returncode != 0:
        print(f"Error compiling Typst PDF 2: {res2.stderr}")
        return
        
    print(f"PDF Successfully compiled to:\n - {PDF_OUT_1}\n - {PDF_OUT_2}")

if __name__ == "__main__":
    main()
