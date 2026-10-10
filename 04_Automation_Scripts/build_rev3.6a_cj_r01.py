#!/usr/bin/env python3
"""
=============================================================================
KALASIN ENGINEERING xCIP-1545 AUTOMATION SYSTEM (Ref: x2608003)
PROJECT SPRINT 18K TPA SPRAY DRYER / JET COOKER - INGREDION (THAILAND) CO., LTD.
CLIENT MASTER I/O AUDIT & MISSING I/O ANALYSIS GENERATOR
OUTPUT FILE: 03-IO_List/Instrument I-O List Rev.3.6a-cj-r01.xlsx
=============================================================================
"""

import os
import shutil
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.abspath(os.path.join(SCRIPT_DIR, ".."))
IO_DIR = os.path.join(BASE_DIR, "03-IO_List")
IO_DIR_ALT = os.path.join(BASE_DIR, "03_IO_Lists_and_Schedules", "IO_List")

INST_FILE = os.path.join(IO_DIR, "Instrument I-O List Rev.3.6a.xlsx")
SLOT_FILE = os.path.join(IO_DIR, "IO_List-By_SlotConfig-rev3.6-r01.xlsx")
OUT_FILE = os.path.join(IO_DIR, "Instrument I-O List Rev.3.6a-cj-r01.xlsx")

# -----------------------------------------------------------------------------
# Color Fills & Styles
# -----------------------------------------------------------------------------
NAVY_FILL = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
SUB_FILL = PatternFill(start_color="2D4A77", end_color="2D4A77", fill_type="solid")
SECTION_FILL = PatternFill(start_color="334E68", end_color="334E68", fill_type="solid")
RED_FILL = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")
AMBER_FILL = PatternFill(start_color="92400E", end_color="92400E", fill_type="solid")

CARD_HEADER_BG = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
CARD_BG_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

ZEBRA_EVEN = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
ZEBRA_ODD = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

MAPPED_FILL = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
MAPPED_FONT = Font(name="Calibri", size=9.5, bold=True, color="166534")

MISSING_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
MISSING_FONT = Font(name="Calibri", size=9.5, bold=True, color="991B1B")

SPARE_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
SPARE_FONT = Font(name="Calibri", size=9.5, bold=True, color="92400E")

FONT_TITLE = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
FONT_SUBTITLE = Font(name="Calibri", size=10, italic=True, color="FFFFFF")
FONT_HEADER = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
FONT_SECTION = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

FONT_DATA = Font(name="Calibri", size=9.5, color="0F172A")
FONT_DATA_BOLD = Font(name="Calibri", size=9.5, bold=True, color="0F172A")
FONT_DATA_CODE = Font(name="Consolas", size=9, color="0F172A")
FONT_DATA_MUTED = Font(name="Calibri", size=9, italic=True, color="64748B")

ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

THIN_BORDER = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)
HEADER_BORDER = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='medium', color='CBD5E1'),
    bottom=Side(style='medium', color='CBD5E1')
)
TOTAL_BORDER = Border(
    top=Side(style='thin', color='94A3B8'),
    bottom=Side(style='double', color='0F172A'),
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1')
)

def set_cell(cell, value, font=None, fill=None, border=THIN_BORDER, alignment=ALIGN_LEFT, num_format=None):
    if not isinstance(cell, openpyxl.cell.cell.MergedCell):
        cell.value = value
    if font: cell.font = font
    if fill: cell.fill = fill
    if border: cell.border = border
    if alignment: cell.alignment = alignment
    if num_format: cell.number_format = num_format

def load_data():
    print(f"Loading Master Instrument List: {INST_FILE}")
    wb_inst = openpyxl.load_workbook(INST_FILE, data_only=True)
    ws_inst = wb_inst['Rev.3']

    print(f"Loading Slot Configuration: {SLOT_FILE}")
    wb_slot = openpyxl.load_workbook(SLOT_FILE, data_only=True)
    slot_sheets = [s for s in wb_slot.sheetnames if s not in ['00_Slot_Index', 'C1SlotConfig', 'C2SlotConfig', 'C3SlotConfig', 'C4SlotConfig', 'C5SlotConfig']]

    # Load slot cards info
    slot_cards = {}
    if '00_Slot_Index' in wb_slot.sheetnames:
        ws_idx = wb_slot['00_Slot_Index']
        for r in range(14, ws_idx.max_row + 1):
            s_tag = ws_idx.cell(r, 1).value
            card = ws_idx.cell(r, 4).value
            if s_tag:
                slot_cards[str(s_tag).strip()] = str(card).strip() if card else ''

    mapped_channels = {}
    for sname in slot_sheets:
        ws = wb_slot[sname]
        chassis = sname[:2]
        card_model = slot_cards.get(sname, '')
        for r in range(6, ws.max_row + 1):
            t_no = ws.cell(r, 1).value
            if t_no is None: continue

            t_desc = ws.cell(r, 2).value
            t_num = ws.cell(r, 3).value
            w_plc = ws.cell(r, 4).value
            w_term = ws.cell(r, 5).value
            dest_raw = ws.cell(r, 6).value
            p_tag = ws.cell(r, 7).value
            i_tag = ws.cell(r, 8).value
            i_desc = ws.cell(r, 9).value
            status = str(ws.cell(r, 10).value or '').strip().upper()

            t_str = str(i_tag or '').strip()
            if t_str and t_str not in ['Spare', 'None', '-']:
                u_tag = t_str.upper()
                u_clean = u_tag.replace(' ', '').replace('-', '')

                info = {
                    'chassis': chassis,
                    'slot': sname,
                    'card_model': card_model,
                    'term_no': t_no,
                    'term_desc': str(t_desc or '').strip(),
                    'term_num': str(t_num or '').strip(),
                    'wire_plc': str(w_plc or '').strip(),
                    'wire_term': str(w_term or '').strip(),
                    'dest': str(dest_raw or '').strip(),
                    'plc_tag': str(p_tag or '').strip(),
                    'inst_tag': t_str,
                    'inst_desc': str(i_desc or '').strip(),
                    'status': status
                }
                if u_tag not in mapped_channels: mapped_channels[u_tag] = []
                mapped_channels[u_tag].append(info)
                if u_clean not in mapped_channels: mapped_channels[u_clean] = []
                mapped_channels[u_clean].append(info)

    all_instruments = []
    for r in range(12, ws_inst.max_row + 1):
        item_no = ws_inst.cell(r, 1).value
        tag = ws_inst.cell(r, 2).value
        new_tag = ws_inst.cell(r, 3).value
        pid = ws_inst.cell(r, 4).value
        desc = ws_inst.cell(r, 5).value
        inst_name = ws_inst.cell(r, 6).value
        do_sig = ws_inst.cell(r, 7).value
        di_sig = ws_inst.cell(r, 8).value
        ao_sig = ws_inst.cell(r, 9).value
        ai_sig = ws_inst.cell(r, 10).value
        bus_sig = ws_inst.cell(r, 11).value
        sig_type = ws_inst.cell(r, 12).value
        sig_to = ws_inst.cell(r, 13).value
        rng = ws_inst.cell(r, 14).value
        fn = ws_inst.cell(r, 15).value
        cable = ws_inst.cell(r, 16).value
        prot = ws_inst.cell(r, 17).value
        remark = ws_inst.cell(r, 18).value

        act_tag = str(new_tag or tag or '').strip()

        if act_tag and act_tag != 'None' and act_tag != 'MAIN PLC/ SCADA':
            t1 = act_tag.upper()
            t2 = str(tag or '').upper()
            t3 = str(new_tag or '').upper()
            t1_c = t1.replace(' ', '').replace('-', '')
            t2_c = t2.replace(' ', '').replace('-', '')
            t3_c = t3.replace(' ', '').replace('-', '')

            hw_matches = (mapped_channels.get(t1) or mapped_channels.get(t1_c) or
                          mapped_channels.get(t2) or mapped_channels.get(t2_c) or
                          mapped_channels.get(t3) or mapped_channels.get(t3_c))

            req_signals = []
            if str(do_sig).strip() in ['1', '1.0'] or (isinstance(do_sig, (int, float)) and do_sig > 0): req_signals.append('DO')
            if str(di_sig).strip() in ['1', '1.0'] or (isinstance(di_sig, (int, float)) and di_sig > 0): req_signals.append('DI')
            if str(ao_sig).strip() in ['1', '1.0'] or (isinstance(ao_sig, (int, float)) and ao_sig > 0): req_signals.append('AO')
            if str(ai_sig).strip() in ['1', '1.0'] or (isinstance(ai_sig, (int, float)) and ai_sig > 0): req_signals.append('AI')
            if str(bus_sig).strip() in ['1', '1.0'] or (isinstance(bus_sig, (int, float)) and bus_sig > 0): req_signals.append('BUS')

            status = 'MAPPED' if hw_matches else 'MISSING'

            prefix = act_tag.split('-')[0] if '-' in act_tag else act_tag

            # Infer missing reason
            if status == 'MISSING':
                if not req_signals or req_signals == ['BUS']:
                    reason = "Local Mechanical Device / Bus Communication"
                    action = "Verify if hardwired I/O or RS485/EtherNet/IP bus interface required"
                elif any(p in prefix for p in ['HS', 'XS', 'XA', 'ES', 'SSL', 'ZSO', 'ZSC']):
                    reason = "Unassigned Digital Pushbutton / Switch / Interlock"
                    action = "Assign spare DI/DO channel on Panel P1-P4 marshaling terminal block"
                elif prefix in ['RVM', 'PCM', 'SCM', 'BF']:
                    reason = "MCC Motor Starter / Drive Control Signal"
                    action = "Route to MCC Control Cabinet or hardwired VFD interlock"
                elif prefix in ['TCV', 'PCV', 'FCV', 'XV', 'ISV', 'BV']:
                    reason = "Unmapped Control / Solenoid Valve"
                    action = "Allocate DO channel for solenoid and AI/AO for positioner"
                else:
                    reason = "Field Process Instrument missing from Hardware Slot Config"
                    action = "Assign available Spare AI/DI channel in nearest Remote I/O Slot"
            else:
                reason = "Hardware Assigned"
                action = "Channel Connected in ControlLogix 1756 PLC Config"

            item = {
                'row': r,
                'item_no': str(item_no).strip() if item_no is not None else '',
                'tag': str(tag).strip() if tag else '',
                'new_tag': str(new_tag).strip() if new_tag else '',
                'active_tag': act_tag,
                'prefix': prefix,
                'pid': str(pid).strip() if pid else '',
                'desc': str(desc).strip() if desc else '',
                'inst_name': str(inst_name).strip() if inst_name else '',
                'req_signals': ', '.join(req_signals) if req_signals else 'LOCAL / NONE',
                'sig_type': str(sig_type).strip() if sig_type else '',
                'sig_to': str(sig_to).strip() if sig_to else '',
                'range': str(rng).strip() if rng else '',
                'cable': str(cable).strip() if cable else '',
                'prot': str(prot).strip() if prot else '',
                'remark': str(remark).strip() if remark else '',
                'status': status,
                'reason': reason,
                'action': action,
                'hw_matches': hw_matches if hw_matches else []
            }
            all_instruments.append(item)

    print(f"Total Master Instruments evaluated: {len(all_instruments)}")
    return all_instruments

def build_all_io_sheet(ws, all_instruments):
    ws.views.sheetView[0].showGridLines = True

    # Title Banner
    ws.merge_cells("A1:Q1")
    set_cell(ws["A1"], "KALASIN ENGINEERING  |  ALL MASTER INSTRUMENT I/O LIST CROSS-REFERENCED TO PLC SLOT CONFIGURATION",
             font=FONT_TITLE, fill=NAVY_FILL, alignment=ALIGN_CENTER)
    ws.row_dimensions[1].height = 25

    ws.merge_cells("A2:Q2")
    set_cell(ws["A2"], "Client Master Instrument Schedule (Rev.3.6a) Mapped to ControlLogix 1756 Standardized PLC Hardware Slots (Rev.3.6a-R01)",
             font=FONT_SUBTITLE, fill=SUB_FILL, alignment=ALIGN_CENTER)
    ws.row_dimensions[2].height = 18

    headers = [
        "Item No.", "Original Tag", "New Tag", "Active Working Tag", "P&ID No.", "Description",
        "Instrument Name", "Signal Req.", "Signal Range", "Protection", "PLC Status",
        "PLC Chassis", "PLC Slot", "Module Catalog", "Terminal Pin", "PLC Wire Tag", "Destination Area"
    ]
    ws.row_dimensions[4].height = 26
    for c_idx, h in enumerate(headers, 1):
        set_cell(ws.cell(4, c_idx), h, font=FONT_HEADER, fill=NAVY_FILL, border=HEADER_BORDER, alignment=ALIGN_CENTER)

    r_idx = 5
    for item in all_instruments:
        fill = ZEBRA_ODD if (r_idx % 2 == 1) else ZEBRA_EVEN
        ws.row_dimensions[r_idx].height = 20

        # If mapped, pull first HW match info
        hw = item['hw_matches'][0] if item['hw_matches'] else {}
        chassis = hw.get('chassis', '-')
        slot = hw.get('slot', '-')
        card = hw.get('card_model', '-')
        t_num = hw.get('term_num', '-')
        w_plc = hw.get('wire_plc', '-')
        dest = hw.get('dest', '-')

        set_cell(ws.cell(r_idx, 1), item['item_no'], font=FONT_DATA, fill=fill, alignment=ALIGN_CENTER)
        set_cell(ws.cell(r_idx, 2), item['tag'], font=FONT_DATA, fill=fill, alignment=ALIGN_LEFT)
        set_cell(ws.cell(r_idx, 3), item['new_tag'], font=FONT_DATA, fill=fill, alignment=ALIGN_LEFT)
        set_cell(ws.cell(r_idx, 4), item['active_tag'], font=FONT_DATA_BOLD, fill=fill, alignment=ALIGN_LEFT)
        set_cell(ws.cell(r_idx, 5), item['pid'], font=FONT_DATA_CODE, fill=fill, alignment=ALIGN_LEFT)
        set_cell(ws.cell(r_idx, 6), item['desc'], font=FONT_DATA, fill=fill, alignment=ALIGN_LEFT)
        set_cell(ws.cell(r_idx, 7), item['inst_name'], font=FONT_DATA, fill=fill, alignment=ALIGN_LEFT)
        set_cell(ws.cell(r_idx, 8), item['req_signals'], font=FONT_DATA, fill=fill, alignment=ALIGN_CENTER)
        set_cell(ws.cell(r_idx, 9), item['range'], font=FONT_DATA, fill=fill, alignment=ALIGN_LEFT)
        set_cell(ws.cell(r_idx, 10), item['prot'], font=FONT_DATA, fill=fill, alignment=ALIGN_CENTER)

        # Status badge
        st = item['status']
        st_fill = MAPPED_FILL if st == 'MAPPED' else MISSING_FILL
        st_font = MAPPED_FONT if st == 'MAPPED' else MISSING_FONT
        set_cell(ws.cell(r_idx, 11), st, font=st_font, fill=st_fill, alignment=ALIGN_CENTER)

        set_cell(ws.cell(r_idx, 12), chassis, font=FONT_DATA_BOLD, fill=fill, alignment=ALIGN_CENTER)
        set_cell(ws.cell(r_idx, 13), slot, font=FONT_DATA_BOLD, fill=fill, alignment=ALIGN_CENTER)
        set_cell(ws.cell(r_idx, 14), card, font=FONT_DATA, fill=fill, alignment=ALIGN_LEFT)
        set_cell(ws.cell(r_idx, 15), t_num, font=FONT_DATA_CODE, fill=fill, alignment=ALIGN_LEFT)
        set_cell(ws.cell(r_idx, 16), w_plc, font=FONT_DATA_CODE, fill=fill, alignment=ALIGN_LEFT)
        set_cell(ws.cell(r_idx, 17), dest, font=FONT_DATA, fill=fill, alignment=ALIGN_LEFT)

        r_idx += 1

    # Auto-fit column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

def build_missing_io_sheet(ws, all_instruments):
    ws.views.sheetView[0].showGridLines = True

    # Title Banner
    ws.merge_cells("A1:K1")
    set_cell(ws["A1"], "KALASIN ENGINEERING  |  MISSING / UNMAPPED I/O AUDIT REPORT (REV.3.6a vs PLC SLOT CONFIG)",
             font=FONT_TITLE, fill=RED_FILL, alignment=ALIGN_CENTER)
    ws.row_dimensions[1].height = 25

    ws.merge_cells("A2:K2")
    set_cell(ws["A2"], "Instruments & Signals from Client Master Schedule (Rev.3.6a) NOT Currently Assigned to ControlLogix 1756 PLC Hardware Slots",
             font=FONT_SUBTITLE, fill=SUB_FILL, alignment=ALIGN_CENTER)
    ws.row_dimensions[2].height = 18

    # KPI Summary Cards (Rows 4-6)
    missing_list = [item for item in all_instruments if item['status'] == 'MISSING']
    mapped_list = [item for item in all_instruments if item['status'] == 'MAPPED']

    tot_count = len(all_instruments)
    mapped_count = len(mapped_list)
    missing_count = len(missing_list)

    mapped_pct = (mapped_count / tot_count * 100.0) if tot_count > 0 else 0.0
    missing_pct = (missing_count / tot_count * 100.0) if tot_count > 0 else 0.0

    cards = [
        ("TOTAL MASTER INSTRUMENTS", f"{tot_count:,}", "Client Rev.3.6a Master Schedule", "A", "C"),
        ("MAPPED TO PLC CONFIG", f"{mapped_count:,}", f"Compliance Rate: {mapped_pct:.1f}%", "D", "F"),
        ("MISSING / UNMAPPED I/O", f"{missing_count:,}", f"Unassigned Rate: {missing_pct:.1f}%", "G", "I"),
        ("AUDIT ACTION REQUIRED", "HIGH PRIORITY", "Review Unassigned Tags below", "J", "K")
    ]

    ws.row_dimensions[4].height = 16
    ws.row_dimensions[5].height = 26
    ws.row_dimensions[6].height = 16

    for title, val, sub, s_col, e_col in cards:
        ws.merge_cells(f"{s_col}4:{e_col}4")
        ws.merge_cells(f"{s_col}5:{e_col}5")
        ws.merge_cells(f"{s_col}6:{e_col}6")

        first_c = ws[f"{s_col}4"]
        c_fill = RED_FILL if "MISSING" in title else (NAVY_FILL if "TOTAL" in title else (MAPPED_FILL if "MAPPED" in title else AMBER_FILL))
        c_head_font = FONT_HEADER if "MAPPED" not in title else Font(name="Calibri", size=10, bold=True, color="166534")

        set_cell(ws[f"{s_col}4"], title, font=c_head_font, fill=c_fill, alignment=ALIGN_CENTER)
        set_cell(ws[f"{s_col}5"], val, font=Font(name="Calibri", size=16, bold=True, color="991B1B" if "MISSING" in title else "0F172A"),
                 fill=CARD_BG_FILL, alignment=ALIGN_CENTER)
        set_cell(ws[f"{s_col}6"], sub, font=Font(name="Calibri", size=8.5, italic=True, color="64748B"),
                 fill=CARD_BG_FILL, alignment=ALIGN_CENTER)

    r_idx = 8
    ws.merge_cells(f"A{r_idx}:K{r_idx}")
    set_cell(ws.cell(r_idx, 1), "DETAILED SCHEDULE OF MISSING / UNMAPPED I/O INSTRUMENTS & SIGNALS",
             font=FONT_SECTION, fill=SECTION_FILL, alignment=ALIGN_LEFT)
    ws.row_dimensions[r_idx].height = 22

    r_idx += 1
    headers = [
        "Item No.", "Tag No.", "New Tag No.", "Active Working Tag", "P&ID No.", "Description",
        "Instrument / Equipment Name", "Required Signal", "Signal To (Client)", "Missing Category / Technical Cause", "Engineering Action Required"
    ]
    ws.row_dimensions[r_idx].height = 26
    for c_idx, h in enumerate(headers, 1):
        set_cell(ws.cell(r_idx, c_idx), h, font=FONT_HEADER, fill=NAVY_FILL, border=HEADER_BORDER, alignment=ALIGN_CENTER)

    r_idx += 1
    for item in missing_list:
        fill = ZEBRA_ODD if (r_idx % 2 == 1) else ZEBRA_EVEN
        ws.row_dimensions[r_idx].height = 20

        set_cell(ws.cell(r_idx, 1), item['item_no'], font=FONT_DATA, fill=fill, alignment=ALIGN_CENTER)
        set_cell(ws.cell(r_idx, 2), item['tag'], font=FONT_DATA, fill=fill, alignment=ALIGN_LEFT)
        set_cell(ws.cell(r_idx, 3), item['new_tag'], font=FONT_DATA, fill=fill, alignment=ALIGN_LEFT)
        set_cell(ws.cell(r_idx, 4), item['active_tag'], font=MISSING_FONT, fill=MISSING_FILL, alignment=ALIGN_LEFT)
        set_cell(ws.cell(r_idx, 5), item['pid'], font=FONT_DATA_CODE, fill=fill, alignment=ALIGN_LEFT)
        set_cell(ws.cell(r_idx, 6), item['desc'], font=FONT_DATA, fill=fill, alignment=ALIGN_LEFT)
        set_cell(ws.cell(r_idx, 7), item['inst_name'], font=FONT_DATA, fill=fill, alignment=ALIGN_LEFT)
        set_cell(ws.cell(r_idx, 8), item['req_signals'], font=FONT_DATA_BOLD, fill=fill, alignment=ALIGN_CENTER)
        set_cell(ws.cell(r_idx, 9), item['sig_to'], font=FONT_DATA, fill=fill, alignment=ALIGN_LEFT)
        set_cell(ws.cell(r_idx, 10), item['reason'], font=FONT_DATA, fill=fill, alignment=ALIGN_LEFT)
        set_cell(ws.cell(r_idx, 11), item['action'], font=FONT_DATA_MUTED, fill=fill, alignment=ALIGN_LEFT)

        r_idx += 1

    # Auto-fit column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

def main():
    print("=================================================================")
    print("BUILDING REVISION WORKBOOK: Instrument I-O List Rev.3.6a-cj-r01.xlsx")
    print("=================================================================")

    all_instruments = load_data()

    print(f"\nOpening Master Workbook template: {INST_FILE}")
    wb = openpyxl.load_workbook(INST_FILE)

    # Sheet 1: All IO List by PLC Config
    if "All IO List by PLC Config" in wb.sheetnames:
        del wb["All IO List by PLC Config"]
    ws_all = wb.create_sheet(title="All IO List by PLC Config")
    print("Generating Sheet: All IO List by PLC Config...")
    build_all_io_sheet(ws_all, all_instruments)

    # Sheet 2: Missing IO List
    if "Missing IO List" in wb.sheetnames:
        del wb["Missing IO List"]
    ws_missing = wb.create_sheet(title="Missing IO List")
    print("Generating Sheet: Missing IO List...")
    build_missing_io_sheet(ws_missing, all_instruments)

    wb.save(OUT_FILE)
    print(f"\nSuccessfully saved primary workbook to: {OUT_FILE}")

    if os.path.exists(IO_DIR_ALT):
        alt_dest = os.path.join(IO_DIR_ALT, os.path.basename(OUT_FILE))
        shutil.copyfile(OUT_FILE, alt_dest)
        print(f"Successfully copied workbook to: {alt_dest}")

    print("=================================================================")
    print("BUILD COMPLETE SUCCESSFUL!")
    print("=================================================================")

if __name__ == '__main__':
    main()
