#!/usr/bin/env python3
"""
=============================================================================
KALASIN ENGINEERING xCIP-1545 AUTOMATION SYSTEM (Ref: x2608003)
PROJECT SPRINT 18K TPA SPRAY DRYER / JET COOKER - INGREDION (THAILAND) CO., LTD.
UPDATE SCRIPT FOR 03-IO_List/xDev_IO_Config/IO_Config.xlsx
=============================================================================
Rules:
1. Panel: [P1, P2, P3, P4, RIO200] checked from 03-IO_List/IO_List-By_SlotConfig_rev02.xlsx
2. IO_Map: just only CxSx With I/O Description (IN-0, Out-0, AI-In-0, AO-Out-0)
   e.g., C1S4-IN-16, C1S8-Out-0, C2S10-AI-In-0, C3S12-AO-Out-0
3. Junction Box: Sample JB401, IS-JB608
4. If New Tag are empty: fill with No Tag-P&ID NO. with running number 1,2,3
=============================================================================
"""

import os
import re
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.abspath(os.path.join(SCRIPT_DIR, ".."))
IO_DIR = os.path.join(BASE_DIR, "03-IO_List")

SLOT_FILE = os.path.join(IO_DIR, "IO_List-By_SlotConfig_rev02.xlsx")
INST_FILE = os.path.join(IO_DIR, "Instrument I-O List Rev.3.6a.xlsx")
CFG_FILE = os.path.join(IO_DIR, "xDev_IO_Config", "IO_Config.xlsx")

def clean_jb(dest_raw):
    dest_clean = str(dest_raw or '').strip()
    if not dest_clean or dest_clean in ['-', 'None']:
        return None
    if dest_clean.startswith('IS-JB-'):
        return 'IS-JB' + dest_clean[6:]
    if dest_clean.startswith('IS-JB'):
        return dest_clean
    if dest_clean.startswith('JB-'):
        return 'JB' + dest_clean[3:]
    if dest_clean.startswith('JB'):
        return dest_clean
    if dest_clean.startswith('RIO-'):
        return 'RIO' + dest_clean[4:]
    if 'CA1' in dest_clean:
        return 'CA1'
    if 'MCC' in dest_clean:
        return 'MCC'
    return dest_clean

def get_io_desc(sname, t_desc, slot_cards):
    card_info = slot_cards.get(sname, {})
    card = card_info.get('card', '')
    fam = card_info.get('fam', '')
    t_desc_str = str(t_desc or '').strip()

    # AI Module (1756-IF16) -> AI-In-0
    if '1756-IF16' in card or 'AI' in fam:
        m = re.search(r'(\d+)', t_desc_str)
        if m:
            return f"AI-In-{m.group(1)}"
        return f"AI-In-{t_desc_str}"
    # AO Module (1756-OF8) -> AO-Out-0
    elif '1756-OF8' in card or 'AO' in fam:
        m = re.search(r'(\d+)', t_desc_str)
        if m:
            return f"AO-Out-{m.group(1)}"
        return f"AO-Out-{t_desc_str}"
    # DI Module (1756-IB32) -> IN-0
    elif '1756-IB32' in card or 'DI' in fam:
        m = re.search(r'(\d+)', t_desc_str)
        if m:
            return f"IN-{m.group(1)}"
        return t_desc_str
    # DO Module (1756-OB32) -> Out-0
    elif '1756-OB32' in card or 'DO' in fam:
        m = re.search(r'(\d+)', t_desc_str)
        if m:
            return f"Out-{m.group(1)}"
        return t_desc_str.replace('OUT-', 'Out-')
    else:
        return t_desc_str

def load_slot_data():
    print(f"Loading Slot Configurations from: {SLOT_FILE}")
    wb_slot = openpyxl.load_workbook(SLOT_FILE, data_only=True)
    slot_sheets = [s for s in wb_slot.sheetnames if s not in ['00_Slot_Index', 'C1SlotConfig', 'C2SlotConfig', 'C3SlotConfig', 'C4SlotConfig', 'C5SlotConfig']]

    # Load module card metadata from 00_Slot_Index
    ws_idx = wb_slot['00_Slot_Index']
    slot_cards = {}
    for r in range(14, ws_idx.max_row + 1):
        s_tag = ws_idx.cell(r, 1).value
        card = ws_idx.cell(r, 4).value
        fam = ws_idx.cell(r, 5).value
        if s_tag:
            slot_cards[str(s_tag).strip()] = {
                'card': str(card).strip() if card else '',
                'fam': str(fam).strip() if fam else ''
            }

    channels_by_tag = {}

    for sname in slot_sheets:
        ws = wb_slot[sname]
        chassis = sname[:2]
        for r in range(6, ws.max_row + 1):
            t_no = ws.cell(r, 1).value
            if t_no is None: continue

            t_desc = str(ws.cell(r, 2).value or '').strip()
            dest_raw = ws.cell(r, 6).value
            i_tag = ws.cell(r, 8).value

            t_str = str(i_tag or '').strip()
            if t_str and t_str not in ['Spare', 'None', '-']:
                u_tag = t_str.upper()
                u_clean = u_tag.replace(' ', '').replace('-', '')

                # Panel [P1,P2,P3,P4,RIO200]
                if chassis == 'C1': panel = 'P1'
                elif chassis == 'C2': panel = 'P2'
                elif chassis == 'C3': panel = 'P3'
                elif chassis == 'C4': panel = 'P4'
                elif chassis == 'C5': panel = 'RIO200'
                else: panel = ''

                jb = clean_jb(dest_raw)

                # IO_Map format: CxSx With I/O Description (IN-0, Out-0, AI-In-0, AO-Out-0)
                # e.g., C1S4-IN-16, C1S8-Out-0, C2S10-AI-In-0, C3S12-AO-Out-0
                io_desc = get_io_desc(sname, t_desc, slot_cards)
                iomap = f"{sname}-{io_desc}"

                info = {
                    'panel': panel,
                    'jb': jb,
                    'iomap': iomap,
                    'slot': sname,
                    'pin': t_no,
                    'pin_desc': t_desc,
                    'io_desc': io_desc
                }

                for k in [u_tag, u_clean]:
                    if k not in channels_by_tag:
                        channels_by_tag[k] = []
                    channels_by_tag[k].append(info)

    print(f"Loaded {len(channels_by_tag)} tag lookup keys from Slot Configuration.")
    return channels_by_tag

def load_master_inst_lookup():
    print(f"Loading Master Instrument List from: {INST_FILE}")
    wb_inst = openpyxl.load_workbook(INST_FILE, data_only=True)
    ws_inst = wb_inst['Rev.3']
    item_lookup = {}
    for r in range(12, ws_inst.max_row + 1):
        item = str(ws_inst.cell(r, 1).value or '').strip()
        tag = ws_inst.cell(r, 2).value
        new_tag = ws_inst.cell(r, 3).value
        if item:
            item_lookup[item] = {
                'tag': str(tag).strip() if tag else '',
                'new_tag': str(new_tag).strip() if new_tag else ''
            }
    print(f"Loaded {len(item_lookup)} items from Rev.3.")
    return item_lookup

def update_io_config():
    channels_by_tag = load_slot_data()
    item_lookup = load_master_inst_lookup()

    print(f"Opening Target Workbook: {CFG_FILE}")
    wb_cfg = openpyxl.load_workbook(CFG_FILE)
    ws = wb_cfg['Sheet2']

    # Running number per P&ID for empty tags
    pid_counter = {}

    filled_tags_count = 0
    filled_hardware_count = 0

    font_data = Font(name="Calibri", size=11, bold=False)
    font_bold = Font(name="Calibri", size=11, bold=True)
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    for r in range(7, ws.max_row + 1):
        item = str(ws.cell(r, 1).value or '').strip()
        new_tag = ws.cell(r, 6).value
        pid = ws.cell(r, 7).value

        # Step 1: Check if New Tag is empty, fill with No Tag-P&ID NO. with running number 1,2,3
        new_tag_str = str(new_tag).strip() if new_tag is not None else ''
        was_empty = False
        if not new_tag_str:
            was_empty = True
            filled_tags_count += 1
            pid_str = str(pid).strip() if pid and str(pid).strip() not in ['None', ''] else ''
            if pid_str:
                pid_counter[pid_str] = pid_counter.get(pid_str, 0) + 1
                new_tag_str = f"No Tag-{pid_str}-{pid_counter[pid_str]}"
            else:
                pid_counter['NO_PID'] = pid_counter.get('NO_PID', 0) + 1
                new_tag_str = f"No Tag-{pid_counter['NO_PID']}"

            ws.cell(r, 6).value = new_tag_str
            ws.cell(r, 6).font = font_data
            ws.cell(r, 6).alignment = align_left

        # Step 2: Match to Slot Config
        # First match by New Tag
        u_new = new_tag_str.upper()
        u_new_c = u_new.replace(' ', '').replace('-', '')
        chs = channels_by_tag.get(u_new) or channels_by_tag.get(u_new_c)

        # Fallback to original Tag from Rev.3.6a
        if not chs and item in item_lookup:
            orig_tag = item_lookup[item]['tag']
            if orig_tag:
                u_orig = orig_tag.upper()
                u_orig_c = u_orig.replace(' ', '').replace('-', '')
                chs = channels_by_tag.get(u_orig) or channels_by_tag.get(u_orig_c)

        if chs:
            filled_hardware_count += 1

            # 1. Panel: [P1,P2,P3,P4,RIO200]
            panel_val = chs[0]['panel']
            ws.cell(r, 3).value = panel_val
            ws.cell(r, 3).font = font_bold
            ws.cell(r, 3).alignment = align_center

            # 2. IO_Map: format CxSx With I/O Description (IN-0, Out-0, AI-In-0, AO-Out-0)
            # Deduplicate entries (e.g. AI channels where IN and i RTN both map to AI-In-0)
            unique_iomaps = []
            for c in chs:
                if c['iomap'] not in unique_iomaps:
                    unique_iomaps.append(c['iomap'])
            iomap_val = ', '.join(unique_iomaps) if unique_iomaps else None
            ws.cell(r, 4).value = iomap_val
            ws.cell(r, 4).font = font_data
            ws.cell(r, 4).alignment = align_left

            # 3. Junction Box: Sample JB401, IS-JB608
            jbs = []
            for c in chs:
                if c['jb'] and c['jb'] not in jbs:
                    jbs.append(c['jb'])
            jb_val = ', '.join(jbs) if jbs else None
            ws.cell(r, 5).value = jb_val
            ws.cell(r, 5).font = font_data
            ws.cell(r, 5).alignment = align_center
        else:
            ws.cell(r, 3).value = None
            ws.cell(r, 4).value = None
            ws.cell(r, 5).value = None

    # Auto-adjust column widths for modified columns
    for col_idx in [3, 4, 5, 6]:
        col_letter = get_column_letter(col_idx)
        max_len = max(len(str(ws.cell(r, col_idx).value or '')) for r in range(6, ws.max_row + 1))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    wb_cfg.save(CFG_FILE)
    print(f"\nSuccessfully updated and saved: {CFG_FILE}")
    print(f"  - Empty tags filled: {filled_tags_count} rows")
    print(f"  - Panel, IO_Map, Junction Box filled: {filled_hardware_count} rows")

if __name__ == '__main__':
    update_io_config()
